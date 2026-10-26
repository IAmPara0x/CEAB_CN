#!/usr/bin/env python3
"""Builds presentation/Project_Presentation.pptx (10 slides, rubric order) from the results."""
import csv
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "report"))
FIG = os.path.join(ROOT, "figures")
OUT = os.path.join(ROOT, "presentation", "Project_Presentation.pptx")

TEAM = [("24BDS0256", "SAHIL BARNWAL"), ("23BCE0298", "Ssunvya jain"),
        ("24BDS0233", "Rishika Singal")]

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
ICE = RGBColor(0xE8, 0xF1, 0xFC)
SAND = RGBColor(0xFD, 0xEE, 0xE7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xB9, 0xCF, 0xEE)
HEAD, BODY = "Cambria", "Calibri"

SUM = list(csv.DictReader(open(os.path.join(ROOT, "results", "summary.csv"))))


def V(sc, param, st, metric, sd=False):
    for r in SUM:
        if r["scenario"] == sc and r["param"] == str(param) and r["strategy"] == st:
            return float(r[f"{metric}_{'sd' if sd else 'mean'}"])
    raise KeyError((sc, param, st, metric))


def pct(new, old, lower=False):
    return (old - new) / old * 100 if lower else (new - old) / old * 100


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = 13.333, 7.5


def bg(slide, color):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = color


def text(slide, x, y, w, h, runs, size=16, color=INK, font=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, name=None, spacing=None):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, {opts})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    paras = runs if isinstance(runs, list) else [runs]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        if spacing:
            para.space_after = Pt(spacing)
        segs = p if isinstance(p, list) else [(p, {})]
        for seg, o in segs:
            r = para.add_run()
            r.text = seg
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.italic = o.get("italic", False)
            r.font.name = o.get("font", font)
            r.font.color.rgb = o.get("color", color)
    return tb


def bullets(slide, x, y, w, h, items, size=16, color=INK, gap=8):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Emu(Inches(0.25))))
        pPr.set("indent", str(-Emu(Inches(0.25))))
        bu = pPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}buChar", {"char": "•"})
        pPr.append(bu)
        segs = it if isinstance(it, list) else [(it, {})]
        for seg, o in segs:
            r = p.add_run()
            r.text = seg
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", False)
            r.font.name = BODY
            r.font.color.rgb = o.get("color", color)
    return tb


def card(slide, x, y, w, h, fill, name=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.06
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    if name:
        s.name = name
    return s


def circle_num(slide, x, y, d, label, fill=BLUE, fg=WHITE, size=18):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = HEAD
    r.font.color.rgb = fg
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE


def title(slide, t, sub=None, dark=False):
    text(slide, 0.6, 0.45, 12.1, 0.8, t, 34, WHITE if dark else NAVY, HEAD, True, name="Title")
    if sub:
        text(slide, 0.6, 1.2, 12.1, 0.45, sub, 16, SOFT if dark else MUTED, BODY, name="Subtitle")


def footer(slide, n, dark=False):
    text(slide, 0.6, 7.0, 9, 0.3, "CEAB · BCSE308P Computer Networks Lab · Fall 2026-27", 10,
         SOFT if dark else MUTED)
    text(slide, 12.2, 7.0, 0.55, 0.3, str(n), 10, SOFT if dark else MUTED, align=PP_ALIGN.RIGHT)


def picture(slide, fname, x, y, w=None, h=None):
    path = os.path.join(FIG, fname)
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)


def stat(slide, x, y, w, big, small, color=BLUE, fill=ICE, h=1.75):
    card(slide, x, y, w, h, fill)
    text(slide, x + 0.25, y + 0.2, w - 0.5, 0.85, big, 40, color, HEAD, True)
    text(slide, x + 0.25, y + 1.05, w - 0.5, h - 1.15, small, 13, INK)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


# numbers
s3b, s3c = V("S3", 50, "beb", "throughput"), V("S3", 50, "ceab", "throughput")
s2b, s2c = V("S2", 20, "beb", "throughput"), V("S2", 20, "ceab", "throughput")
p99b, p99c = V("S3", 50, "beb", "access_delay_p99"), V("S3", 50, "ceab", "access_delay_p99")
jb, jc = V("S3", 50, "beb", "jain_short"), V("S3", 50, "ceab", "jain_short")
pcb, pcc = V("S3", 50, "beb", "collision_prob"), V("S3", 50, "ceab", "collision_prob")
n100b, n100c = V("sweep_n", 100, "beb", "throughput"), V("sweep_n", 100, "ceab", "throughput")
nsims = sum(int(r["runs"]) for r in SUM)

# crop the CEAB half of the flowchart for slide 5
fc = Image.open(os.path.join(FIG, "fig_flowchart.png"))
fc.crop((int(fc.width * 0.47), 0, fc.width, int(fc.height * 0.80))).save(os.path.join(FIG, "fig_flowchart_ceab.png"))

# ---------------------------------------------------------------- 1. title and team
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
text(s, 0.8, 0.8, 11, 0.4, "BCSE308P COMPUTER NETWORKS LAB · INNOVATIVE EVALUATION PROJECT", 13, SOFT, bold=True)
text(s, 0.8, 1.5, 11.5, 2.0, [[("Contender-Estimation Adaptive Backoff", {})],
                              [("Teaching Wi-Fi stations how long to wait", {"size": 26, "bold": False, "color": SOFT})]],
     44, WHITE, HEAD, True, spacing=10)
text(s, 0.8, 3.65, 11.5, 0.9, "A comparative study of nine CSMA/CA backoff strategies in IEEE 802.11, "
     "and a new scheme that sets the contention window from the number of stations it overhears",
     17, WHITE)
for i, (reg, nm) in enumerate(TEAM):
    x = 0.8 + i * 3.9
    card(s, x, 5.2, 3.6, 1.15, RGBColor(0x16, 0x34, 0x5C))
    text(s, x + 0.25, 5.35, 3.1, 0.45, nm, 18, WHITE, bold=True)
    text(s, x + 0.25, 5.8, 3.1, 0.35, reg, 13, SOFT)
text(s, 0.8, 6.75, 11, 0.35, "School of Computer Science and Engineering · VIT · October 2026", 12, SOFT)
notes(s, "Introduce the team and the one-line idea: the 802.11 backoff rule forgets how crowded the "
         "channel is after every frame; CEAB measures the crowd and waits accordingly.")

# ---------------------------------------------------------------- 2. problem and motivation
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Crowded Wi-Fi loses a quarter of its capacity",
      "Standard 802.11 binary exponential backoff (BEB), 50 saturated stations, 802.11g 54 Mb/s")
stat(s, 0.6, 1.95, 3.85, f"−{pct(s3b, s3c, True):.0f} %", "throughput lost against the best "
     f"backoff at N = 50 ({s3b:.1f} vs {s3c:.1f} Mb/s)", ORANGE, SAND)
stat(s, 4.75, 1.95, 3.85, f"{pcb:.2f}", "probability that a transmission attempt collides", ORANGE, SAND)
stat(s, 8.9, 1.95, 3.85, f"{p99b:.0f} ms", "99th-percentile access delay: one frame in a hundred "
     "waits a third of a second", ORANGE, SAND)
text(s, 0.6, 4.1, 5.9, 0.4, "Why it happens", 20, NAVY, HEAD, True)
bullets(s, 0.6, 4.6, 5.9, 2.3, [
    "After every success BEB resets the window to 16 slots, regardless of how many stations compete",
    "The next frame collides, doubles, collides again: contention is re-learned for every frame",
    "The winner keeps the smallest window and tends to win again (capture effect)",
], 15)
text(s, 6.95, 4.1, 5.8, 0.4, "Why it matters", 20, NAVY, HEAD, True)
bullets(s, 6.95, 4.6, 5.8, 2.3, [
    "Lecture halls, hostels and IoT deployments put dozens of active stations on one access point",
    f"Short-term fairness at N = 50 is only {jb:.2f} (Jain index; 1.0 = perfectly fair)",
    "Tail delay, not the mean, is what breaks calls, games and video",
], 15)
footer(s, 2)
notes(s, "All numbers are from our simulator, 10 seeds each. Emphasise that the problem grows with "
         "the number of stations, which is the trend in real deployments.")

# ---------------------------------------------------------------- 3. existing approach and gap
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Existing waiting strategies either forget or assume", "Baseline: IEEE 802.11 DCF with binary exponential backoff")
card(s, 0.6, 1.9, 3.9, 4.8, SAND)
text(s, 0.85, 2.1, 3.4, 0.4, "Baseline: BEB", 20, ORANGE, HEAD, True)
steps = ["Draw backoff b ~ U[0, W−1]", "Count down idle 9 µs slots, freeze while busy",
         "Transmit at b = 0; wait for the ACK", "No ACK → W ← min(2W, 1024), retry ≤ 7",
         "ACK → W ← 16 (forget everything)"]
for i, st_ in enumerate(steps):
    circle_num(s, 0.85, 2.7 + i * 0.75, 0.45, str(i + 1), ORANGE, WHITE, 14)
    text(s, 1.45, 2.72 + i * 0.75, 2.9, 0.6, st_, 14, INK)
card(s, 4.8, 1.9, 3.85, 4.8, RGBColor(0xF2, 0xF2, 0xF2))
text(s, 5.05, 2.1, 3.4, 0.4, "History-based rules", 20, NAVY, HEAD, True)
text(s, 5.05, 2.55, 3.4, 0.4, "Linear · Fibonacci · Quadratic · MILD · EIED · Constant", 12, MUTED)
bullets(s, 5.05, 3.1, 3.4, 3.4, [
    "Change W by a fixed rule after each outcome",
    "Never learn how many stations compete",
    "Linear/Fibonacci collide even more than BEB; MILD can lock at CWmax",
], 14)
card(s, 8.95, 1.9, 3.8, 4.8, ICE)
text(s, 9.2, 2.1, 3.4, 0.4, "Estimation-based", 20, NAVY, HEAD, True)
text(s, 9.2, 2.55, 3.4, 0.4, "Calì et al. 2000 · Bianchi & Tinnirello 2003 · Idle Sense 2005", 12, MUTED)
bullets(s, 9.2, 3.1, 3.35, 3.4, [
    "Aim for the optimal window, but model the other stations",
    "Model inversion is fragile: real counters are correlated",
    "Idle Sense next to legacy stations gets 0.9 Mb/s vs 27 Mb/s",
], 14)
footer(s, 3)
notes(s, "Gap: we need an estimate of contention that is model-free, passive, and independent of "
         "what backoff the other stations run.")

# ---------------------------------------------------------------- 4. proposed novelty
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "CEAB counts who is talking, then waits",
      "Contender-Estimation Adaptive Backoff: four steps, run independently by every station")
cards = [
    ("1", "Overhear", "Log the sender address (TA) of the last M = 200 decoded frames",
     "information every 802.11 header already carries"),
    ("2", "Estimate", "n_eff = M(M−1) / Σ c_j(c_j−1)", "inverse Simpson index: equals n for n equal senders"),
    ("3", "Correct for load", "ρ = min(1, I* / I_obs);  n_c = ρ · n_eff", "long idle gaps → fewer backlogged stations"),
    ("4", "Set the window", "W* = 2 / τ*(n_c) − 1", "τ* = throughput-optimal attempt probability (Bianchi)"),
]
for i, (n_, head, body, small) in enumerate(cards):
    x = 0.6 + i * 3.1
    card(s, x, 1.95, 2.85, 3.3, ICE)
    circle_num(s, x + 0.25, 2.2, 0.6, n_, BLUE, WHITE, 20)
    text(s, x + 0.25, 2.95, 2.4, 0.45, head, 20, NAVY, HEAD, True)
    text(s, x + 0.25, 3.45, 2.4, 1.0, body, 15, INK, bold=i in (1, 2, 3))
    text(s, x + 0.25, 4.45, 2.4, 0.7, small, 12, MUTED)
text(s, 0.6, 5.55, 12.1, 0.4, "What exactly differs from the baseline", 20, NAVY, HEAD, True)
bullets(s, 0.6, 6.0, 6.0, 1.0, [[("After success: ", {"bold": True}), ("W stays at the optimum, never back to 16", {})],
                                [("After collision: ", {"bold": True}), ("W unchanged (K = 0), not doubled", {})]], 14, gap=4)
bullets(s, 6.75, 6.0, 6.0, 1.0, [[("Estimator: ", {"bold": True}), ("model-free; valid whatever the others run", {})],
                                 [("Same input for all: ", {"bold": True}), ("all stations hold the same window → fairness", {})]], 14, gap=4)
footer(s, 4)
notes(s, "Novelty: the inverse Simpson index from ecology has not been used for 802.11 backoff "
         "tuning as far as we found. Unlike Calì/Bianchi-Tinnirello, no model of the other "
         "stations; unlike Idle Sense, no fixed idle-slot target, so it does not starve next to "
         "legacy stations.")

# ---------------------------------------------------------------- 5. architecture / algorithm
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Architecture and algorithm", "Event-driven 802.11 DCF simulator with pluggable backoff strategies")
pic = picture(s, "fig_architecture.png", 0.6, 1.85, w=6.6)
pic2 = picture(s, "fig_flowchart_ceab.png", 7.6, 1.75, h=5.1)
text(s, 0.6, 5.75, 6.6, 1.1, [[("Engine: ", {"bold": True}),
                               ("jumps straight to the next slot in which a backoff counter hits zero, so "
                                f"{nsims} ten-second runs take about 6 minutes on 16 cores", {})]], 14)
footer(s, 5)
notes(s, "Left: each station = traffic source + queue + DCF state machine + strategy object. "
         "The channel engine decides success/collision/error and broadcasts idle gap and TA. "
         "Right: CEAB flowchart; estimation runs every busy period, window update on own outcomes.")

# ---------------------------------------------------------------- 6. implementation setup
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Implementation setup", "Python 3.14 · NumPy · SciPy · Matplotlib · 9 unit tests incl. Bianchi validation")
rows = [("PHY", "802.11g ERP-OFDM, 54 Mb/s data, 24 Mb/s ACK"), ("Slot / SIFS / DIFS", "9 / 10 / 28 µs"),
        ("Payload", "1500 B (256–2304 B sweep)"), ("Ts / Tc (basic)", "314 / 323 µs"),
        ("CWmin / CWmax / retries", "15 / 1023 / 7"), ("Stations N", "2 – 100"),
        ("Traffic", "saturated, or Poisson 2 – 40 Mb/s"), ("Runs", "10 seeds × 10 s, 1 s warm-up"),
        ("CEAB", "M = 200, γ = 0.5, β = 0.05, K = 0")]
tbl = s.shapes.add_table(len(rows) + 1, 2, Inches(0.6), Inches(1.9), Inches(6.3), Inches(4.6)).table
tbl.columns[0].width = Inches(2.5)
tbl.columns[1].width = Inches(3.8)
for j, h_ in enumerate(("Parameter", "Value")):
    c = tbl.cell(0, j)
    c.text = h_
    c.fill.solid()
    c.fill.fore_color.rgb = NAVY
    for r in c.text_frame.paragraphs[0].runs:
        r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(14), True, WHITE, BODY
for i, (a, b) in enumerate(rows, 1):
    for j, v in enumerate((a, b)):
        c = tbl.cell(i, j)
        c.text = v
        c.fill.solid()
        c.fill.fore_color.rgb = ICE if i % 2 else WHITE
        for r in c.text_frame.paragraphs[0].runs:
            r.font.size, r.font.color.rgb, r.font.name = Pt(14), INK, BODY
            r.font.bold = j == 0
text(s, 7.4, 1.9, 5.3, 0.4, "Code structure", 20, NAVY, HEAD, True)
mods = [("phy.py", "frame timing, Ts and Tc"), ("analytical.py", "Bianchi model, τ*(n) table"),
        ("strategies.py", "9 strategies incl. CEAB"), ("simulator.py", "event-driven DCF engine"),
        ("experiments.py", f"{nsims} parallel runs"), ("analysis.py", "mean ± SD, t-tests, plots"),
        ("demo.py", "live CLI, trace, GUI")]
for i, (m, d) in enumerate(mods):
    y = 2.45 + i * 0.6
    card(s, 7.4, y, 5.3, 0.5, ICE if i % 2 == 0 else RGBColor(0xF4, 0xF8, 0xFD))
    text(s, 7.6, y + 0.1, 1.9, 0.35, m, 14, BLUE, "Courier New", True)
    text(s, 9.6, y + 0.1, 3.0, 0.35, d, 14, INK)
footer(s, 6)
notes(s, "Validation: simulated BEB matches Bianchi's analytical throughput within 1.7 % for "
         "2–100 stations, and an independent slot-by-slot Monte Carlo in the tests.")

# ---------------------------------------------------------------- 7. scenarios and metrics
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Scenarios and metrics", f"{nsims} independent simulation runs, every point = mean ± SD of 10 seeds")
scen = [("S1 · S2 · S3", "saturated, N = 5 / 20 / 50 (low / medium / high)"),
        ("Station sweep", "N = 2 … 100, all nine strategies"),
        ("Offered load", "Poisson traffic, 30 stations, 2 – 40 Mb/s"),
        ("Robustness", "frame errors 0 – 30 %, RTS/CTS, payload 256 – 2304 B"),
        ("Churn", "5 → 50 → 15 stations during a run"),
        ("Coexistence", "10 legacy BEB + 10 new-scheme stations"),
        ("Ablation", "CEAB escalation cap K and window M")]
for i, (a, b) in enumerate(scen):
    y = 1.95 + i * 0.68
    circle_num(s, 0.6, y, 0.45, str(i + 1), BLUE, WHITE, 14)
    text(s, 1.25, y + 0.03, 1.9, 0.4, a, 15, NAVY, bold=True)
    text(s, 3.15, y + 0.03, 3.6, 0.6, b, 14, INK)
card(s, 7.2, 1.9, 5.55, 4.85, ICE)
text(s, 7.45, 2.1, 5.0, 0.4, "Metrics and why", 20, NAVY, HEAD, True)
bullets(s, 7.45, 2.65, 5.1, 4.0, [
    [("Throughput", {"bold": True}), (" — capacity; what the optimal window maximises", {})],
    [("Collision probability", {"bold": True}), (" — direct sign of a window that is too small", {})],
    [("99th-pct access delay", {"bold": True}), (" — the tail is what real-time traffic feels", {})],
    [("Short-term Jain fairness", {"bold": True}), (" — 100 ms windows expose the capture effect", {})],
    [("Drop rate, end-to-end delay", {"bold": True}), (" — losses and queueing seen by upper layers", {})],
    [("Welch t-test", {"bold": True}), (" — is the difference real, not noise?", {})],
], 14, gap=7)
footer(s, 7)
notes(s, "Mention that every scenario can be re-run live with demo.py by changing a flag or slider.")

# ---------------------------------------------------------------- 8. results
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "CEAB stays near capacity as the network grows",
      "Saturation throughput, all nine strategies, N = 2 … 100 (10 runs per point)")
picture(s, "fig_throughput_vs_n.png", 0.5, 1.8, w=8.2)
stat(s, 9.0, 1.85, 3.75, f"+{pct(s3c, s3b):.1f} %", f"throughput at N = 50 ({s3b:.1f} → {s3c:.1f} Mb/s), "
     f"+{pct(n100c, n100b):.0f} % at N = 100", BLUE, ICE, 1.55)
stat(s, 9.0, 3.6, 3.75, f"−{pct(p99c, p99b, True):.0f} %", f"99th-pct delay at N = 50 "
     f"({p99b:.0f} → {p99c:.0f} ms)", BLUE, ICE, 1.55)
stat(s, 9.0, 5.35, 3.75, f"{jb:.2f} → {jc:.2f}", f"short-term fairness at N = 50; collisions "
     f"{pcb:.2f} → {pcc:.2f}", BLUE, ICE, 1.55)
footer(s, 8)
notes(s, f"All BEB vs CEAB differences in S1–S3 are significant (Welch p < 0.01) except mean "
         f"access delay at N = 50, which is fixed by the frame rate in saturation. S1 gain is "
         f"small (+{pct(V('S1', 5, 'ceab', 'throughput'), V('S1', 5, 'beb', 'throughput')):.1f} %) "
         f"because BEB is already near-optimal with 5 stations; S2 +{pct(s2c, s2b):.1f} %.")

# ---------------------------------------------------------------- 9. analysis and limitations
s = prs.slides.add_slide(BLANK)
bg(s, WHITE)
title(s, "Where it helps, and where it does not", "Dynamic membership trace (left) and critical evaluation (right)")
picture(s, "fig_churn.png", 0.5, 1.75, h=5.15)
card(s, 6.6, 1.8, 6.15, 2.35, ICE)
text(s, 6.85, 1.95, 5.7, 0.4, "Strengths", 18, BLUE, HEAD, True)
bullets(s, 6.85, 2.4, 5.7, 1.7, [
    "Tracks 5 → 50 → 15 stations within ~0.1 s",
    f"Carries {V('load', 40, 'ceab', 'throughput'):.1f} Mb/s of Poisson traffic vs "
    f"{V('load', 40, 'beb', 'throughput'):.1f} for BEB; light-load delay unchanged",
    "Adopters are better off even beside legacy BEB (Idle Sense starves)",
], 13, gap=4)
card(s, 6.6, 4.35, 6.15, 2.45, SAND)
text(s, 6.85, 4.5, 5.7, 0.4, "Limitations and failure cases", 18, ORANGE, HEAD, True)
bullets(s, 6.85, 4.95, 5.7, 1.8, [
    f"RTS/CTS, N = 5: no gain ({pct(V('rts', 5, 'ceab', 'throughput'), V('rts', 5, 'beb', 'throughput')):+.1f} %) — collisions are cheap",
    "FER 30 %: gains vanish, EIED marginally better",
    f"Medium Poisson load: mean delay {V('load', 20, 'ceab', 'e2e_delay_mean'):.2f} vs "
    f"{V('load', 20, 'beb', 'e2e_delay_mean'):.2f} ms",
    "Hidden terminals not modelled; legacy BEB keeps a 2.6 : 1 share",
], 13, gap=4)
footer(s, 9)
notes(s, "Be upfront about the failure cases: this is what the rubric's critical-evaluation "
         "criterion looks for. Ablation: per-retry doubling (K = 1, 2) raised the tail delay, so K = 0.")

# ---------------------------------------------------------------- 10. conclusion
s = prs.slides.add_slide(BLANK)
bg(s, NAVY)
title(s, "Conclusion, future work and references", dark=True)
text(s, 0.6, 1.45, 6.2, 0.4, "Conclusion", 20, WHITE, HEAD, True)
bullets(s, 0.6, 1.95, 6.0, 3.0, [
    "Waiting well = remembering how crowded the channel is",
    f"CEAB: +{pct(s2c, s2b):.0f} % / +{pct(s3c, s3b):.0f} % throughput at N = 20 / 50, "
    f"−{pct(p99c, p99b, True):.0f} % tail delay, fairness {jc:.2f}",
    "Within 2.5 % of the theoretical limit for 5 – 100 stations",
    "Uses only what every 802.11 receiver already decodes",
], 15, WHITE, gap=8)
text(s, 0.6, 4.35, 6.2, 0.4, "Future work", 20, WHITE, HEAD, True)
bullets(s, 0.6, 4.85, 6.0, 1.6, [
    "Hidden terminals and multi-cell; airtime-aware estimate for multi-rate",
    "ns-3 port and an OpenWiFi / ath9k testbed; 802.11ah IoT scale",
], 15, WHITE, gap=6)
card(s, 7.2, 1.45, 5.55, 5.35, RGBColor(0x16, 0x34, 0x5C))
text(s, 7.45, 1.6, 5.0, 0.4, "Key references", 18, WHITE, HEAD, True)
refs = ["[1] IEEE Std 802.11-2020",
        "[2] G. Bianchi, IEEE JSAC 18(3), 2000",
        "[3] F. Calì et al., IEEE/ACM ToN 8(6), 2000",
        "[4] V. Bharghavan et al., MACAW, SIGCOMM 1994",
        "[5] N.-O. Song et al., EIED, IEEE VTC 2003",
        "[6] M. Heusse et al., Idle Sense, SIGCOMM 2005",
        "[7] G. Bianchi, I. Tinnirello, INFOCOM 2003",
        "[16] E. H. Simpson, Nature 163, 1949",
        "Full IEEE-style list: project report, Ch. 7"]
text(s, 7.45, 2.15, 5.1, 4.5, refs, 13, SOFT, spacing=6)
footer(s, 10, dark=True)
notes(s, "Close with the demo offer: python demo.py gui — change N, load, error rate or access mode "
         "live and compare BEB with CEAB.")

prs.save(OUT)
print("saved", OUT)
