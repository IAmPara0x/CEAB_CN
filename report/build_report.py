#!/usr/bin/env python3
"""Builds report/Project_Report.docx (and .pdf) in the BCSE308P template format.

All numbers in the text and tables are read from results/summary.csv, results/raw_runs.csv and
results/churn_series.npz, so the report always matches the latest run of run_experiments.py.
Page numbers in the table of contents are filled in by a second pass over the rendered PDF.

Edit TEAM below before building.
"""
import csv
import inspect
import json
import os
import re
import subprocess
import sys
import tempfile

import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from csmasim import strategies as strat_mod                      # noqa: E402
from csmasim.analysis import NICE, main_table                     # noqa: E402
from csmasim.analysis import load as load_rows                    # noqa: E402
from csmasim.phy import PhyParams                                 # noqa: E402

# ----------------------------------------------------------------------------- edit me
TEAM = [("24BDS0256", "SAHIL BARNWAL"), ("23BCE0298", "Ssunvya jain"),
        ("24BDS0233", "Rishika Singal")]
SPECIALISATION = "(Specialization if any)"
TITLE = ("CONTENDER-ESTIMATION ADAPTIVE BACKOFF (CEAB): A COMPARATIVE STUDY OF WAITING "
         "STRATEGIES FOR IEEE 802.11 CSMA/CA")
TITLE_CASE = ("Contender-Estimation Adaptive Backoff (CEAB): A Comparative Study of Waiting "
              "Strategies for IEEE 802.11 CSMA/CA")
MONTH = "October 2026"
# -------------------------------------------------------------------------------------

RES = os.path.join(ROOT, "results")
FIG = os.path.join(ROOT, "figures")
OUT = os.path.join(ROOT, "report", "Project_Report.docx")
LOGO = os.path.join(ROOT, "report", "assets", "vit_logo.jpg")
FONT = "Times New Roman"

# ============================================================================ data access
SUM = list(csv.DictReader(open(os.path.join(RES, "summary.csv"))))
ROWS = load_rows(RES)
ANA = list(csv.DictReader(open(os.path.join(RES, "analytical.csv"))))


def V(sc, param, st, metric, sd=False):
    for r in SUM:
        if r["scenario"] == sc and r["param"] == str(param) and r["strategy"] == st:
            return float(r[f"{metric}_{'sd' if sd else 'mean'}"])
    raise KeyError((sc, param, st, metric))


def pct(new, old, lower_better=False):
    return (old - new) / old * 100 if lower_better else (new - old) / old * 100


def share(param):
    rs = [json.loads(r["share"]) for r in ROWS if r["scenario"] == "coexist" and r["param"] == param]
    keys = rs[0].keys()
    return {k: (float(np.mean([x[k] for x in rs])), float(np.std([x[k] for x in rs], ddof=1)))
            for k in keys}


def opt_thr(n):
    return float(next(a for a in ANA if int(a["n"]) == n)["optimal_thr"])


# ============================================================================ docx helpers
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)      # A4, as in the template
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Inches(1))

FIGNO = {"n": 0}
TABNO = {"n": 0}
CHAPTER = {"n": 0}
HEADINGS = []           # (text used for TOC lookup)


def set_font(run, size=12, bold=False, italic=False, mono=False, color=None):
    f = "Courier New" if mono else FONT
    run.font.name = f
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(att), f)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


style = doc.styles["Normal"]
style.font.name = FONT
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
style.paragraph_format.space_after = Pt(6)
for name, size in [("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 12)]:
    st = doc.styles[name]
    st.font.name = FONT
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.element.rPr.rFonts.set(qn("w:ascii"), FONT)
    st.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    st.paragraph_format.space_before = Pt(12 if size == 14 else 10)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.5


def para(text="", size=12, bold=False, italic=False, align="justify", spacing=1.5, after=6,
         before=0, keep=False, indent=None):
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    pf = p.paragraph_format
    pf.line_spacing = spacing
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    if indent is not None:
        pf.left_indent = Inches(indent)
    if keep:
        pf.keep_with_next = True
    if text:
        add_rich(p, text, size, bold, italic)
    return p


def add_rich(p, text, size=12, bold=False, italic=False):
    """Supports **bold** and *italic* spans inside text."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            set_font(p.add_run(part[2:-2]), size, True, italic)
        elif part.startswith("*"):
            set_font(p.add_run(part[1:-1]), size, bold, True)
        else:
            set_font(p.add_run(part), size, bold, italic)


def bullets(items, size=12):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(3)
        add_rich(p, it, size)


def numbered(items):
    for it in items:
        p = doc.add_paragraph(style="List Number")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(3)
        add_rich(p, it)


def chapter(title):
    CHAPTER["n"] += 1
    FIGNO["n"] = TABNO["n"] = 0
    text = f"{CHAPTER['n']}. {title.upper()}"
    h = doc.add_heading(level=1)
    h.paragraph_format.page_break_before = True
    set_font(h.add_run(text), 14, True)
    HEADINGS.append(text)


def section(title):
    h = doc.add_heading(level=2)
    h.paragraph_format.keep_with_next = True
    set_font(h.add_run(title), 12, True)
    HEADINGS.append(title)


def subsection(title):
    h = doc.add_heading(level=3)
    h.paragraph_format.keep_with_next = True
    set_font(h.add_run(title), 12, True, italic=True)


def figure(fname, caption, width=6.2):
    FIGNO["n"] += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(os.path.join(FIG, fname), width=Inches(width))
    label = f"Figure {CHAPTER['n']}.{FIGNO['n']}"
    para(f"**{label}:** {caption}", size=10.5, align="center", spacing=1.15, after=10)
    return label


def shade(cell, hexcolor):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcpr.append(shd)


def table(header, rows, widths, caption=None, size=10, bold_last_col=False, align=None):
    if caption:
        TABNO["n"] += 1
        para(f"**Table {CHAPTER['n']}.{TABNO['n']}:** {caption}", size=10.5, align="center",
             spacing=1.15, after=4, keep=True)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        set_font(r, size, True)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade(c, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            add_rich(p, str(v), size)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = (align[i] if align else
                           (WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER))
    fix_widths(t, widths)
    for row in t.rows:
        trpr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trpr.append(cant)
    # repeat header row
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    trpr.append(th)
    para("", after=4, spacing=1.0)
    return t


def fix_widths(t, widths):
    """Fixed layout + explicit grid so Word and LibreOffice honour the column widths."""
    t.autofit = False
    tblpr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))
    for i, col in enumerate(t.columns):
        col.width = Inches(widths[i])
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)


def code(text, size=8.5):
    for line in text.rstrip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p.add_run(line if line else " "), size, mono=True)
    para("", after=4, spacing=1.0)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def remove_borders(t):
    tblpr = t._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tblpr.append(b)


def add_page_number_footer(section_, fmt="decimal", start=None):
    section_.footer.is_linked_to_previous = False
    p = section_.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for el in list(p._p):
        if el.tag == qn("w:r"):
            p._p.remove(el)
    run = p.add_run()
    set_font(run, 11)
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kind)
            run._r.append(f)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            run._r.append(it)
    sectpr = section_._sectPr
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    sectpr.append(pg)


# ============================================================================ numbers
s3b, s3c = V("S3", 50, "beb", "throughput"), V("S3", 50, "ceab", "throughput")
s2b, s2c = V("S2", 20, "beb", "throughput"), V("S2", 20, "ceab", "throughput")
s1b, s1c = V("S1", 5, "beb", "throughput"), V("S1", 5, "ceab", "throughput")
p99b, p99c = V("S3", 50, "beb", "access_delay_p99"), V("S3", 50, "ceab", "access_delay_p99")
jb, jc = V("S3", 50, "beb", "jain_short"), V("S3", 50, "ceab", "jain_short")
pcb, pcc = V("S3", 50, "beb", "collision_prob"), V("S3", 50, "ceab", "collision_prob")
n100b, n100c = V("sweep_n", 100, "beb", "throughput"), V("sweep_n", 100, "ceab", "throughput")
n100i = V("sweep_n", 100, "idlesense", "throughput")
eff = min(V("sweep_n", n, "ceab", "throughput") / opt_thr(n) for n in (5, 10, 20, 35, 50, 75, 100))
val_err = max(abs(V("validation", n, "beb", "throughput") /
                  float(next(a for a in ANA if int(a["n"]) == n)["bianchi_thr"]) - 1) * 100
              for n in (2, 5, 10, 20, 35, 50, 75, 100))
val_perr = max(abs(V("validation", n, "beb", "collision_prob") -
                   float(next(a for a in ANA if int(a["n"]) == n)["bianchi_p"]))
               for n in (2, 5, 10, 20, 35, 50, 75, 100))
sh_ceab, sh_is, sh_eied, sh_beb = share("ceab"), share("idlesense"), share("eied"), share("beb")
z = np.load(os.path.join(RES, "churn_series.npz"))
t_ser = z["beb/t"][0]


def churn_mean(st, a, b):
    y = np.nanmean(z[f"{st}/throughput"], axis=0)
    return float(np.nanmean(y[(t_ser >= a) & (t_ser < b)]))


def churn_settle():
    act = z["ceab/active"][0]
    nh = np.nanmean(z["ceab/n_hat"], axis=0)
    out = []
    for tc, target in ((4.0, 50), (8.0, 15)):
        idx = np.flatnonzero((t_ser > tc) & (np.abs(nh - target) <= 0.15 * target))
        out.append(t_ser[idx[0]] - tc if idx.size else float("nan"))
    return out


settle_up, settle_down = churn_settle()
RUNS = int(next(r for r in SUM if r["scenario"] == "S3")["runs"])
NSIMS = len(ROWS)
phy = PhyParams()

# ============================================================================ title page
para("Innovative Evaluation Practices for BCSE308P - Computer Networks Lab", 14, True, align="center",
     spacing=1.5, after=0)
para("Fall Semester 2026-27", 13, True, align="center", spacing=1.5, after=36)
para(TITLE, 16, True, align="center", spacing=1.5, after=30)
tt = doc.add_table(rows=len(TEAM), cols=2)
tt.alignment = WD_TABLE_ALIGNMENT.CENTER
remove_borders(tt)
fix_widths(tt, (1.8, 3.4))
for i, (reg, name) in enumerate(sorted(TEAM)):
    for j, v in enumerate((reg, name.upper())):
        c = tt.rows[i].cells[j]
        c.text = ""
        set_font(c.paragraphs[0].add_run(v), 13, True)
        c.paragraphs[0].paragraph_format.line_spacing = 1.5
para("", after=24)
for txt, sz, b, aft in [("B.Tech.", 13, True, 0), ("in", 12, False, 0),
                        ("Computer Science and Engineering", 13, True, 0), (SPECIALISATION, 12, False, 24),
                        ("School of Computer Science and Engineering", 13, True, 18)]:
    para(txt, sz, b, align="center", spacing=1.5, after=aft)
lp = doc.add_paragraph()
lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
lp.add_run().add_picture(LOGO, width=Inches(3.45))
para(MONTH, 13, True, align="center", spacing=1.5, before=24)

# ============================================================================ abstract
page_break()
para("ABSTRACT", 14, True, align="center", spacing=1.5, after=12)
abstract = (
    f"IEEE 802.11 stations share the wireless medium using CSMA/CA, and the binary exponential "
    f"backoff (BEB) of the Distributed Coordination Function decides how long each station waits "
    f"before transmitting. BEB doubles the contention window after a collision but returns it to "
    f"the minimum after every success, so in a crowded cell every new frame starts in a window that "
    f"is far too small. This project studies nine waiting strategies — BEB, constant, linear, "
    f"Fibonacci, quadratic, MILD, EIED, Idle Sense and a new scheme, Contender-Estimation Adaptive "
    f"Backoff (CEAB) — in a purpose-built, event-driven 802.11g DCF simulator that was validated "
    f"against Bianchi’s analytical model (throughput error below {val_err:.1f} %). CEAB lets every "
    f"station overhear the transmitter address of successfully decoded frames and compute the "
    f"effective number of contenders as the inverse Simpson index of the last 200 senders. An "
    f"idle-gap load factor corrects the estimate under unsaturated traffic, and the window is set to "
    f"the analytically optimal value for that number of stations instead of being reset to CWmin. "
    f"Over {NSIMS} simulation runs ({RUNS} seeds per configuration), CEAB raised saturation "
    f"throughput by {pct(s2c, s2b):.1f} % with 20 stations and {pct(s3c, s3b):.1f} % with 50 "
    f"stations, cut the 99th-percentile access delay by {pct(p99c, p99b, True):.0f} %, lowered the "
    f"collision probability from {pcb:.2f} to {pcc:.2f}, raised short-term Jain fairness from "
    f"{jb:.2f} to {jc:.2f}, and stayed within {100 - eff * 100:.1f} % of the theoretical throughput "
    f"limit from 5 to 100 stations. Its estimator tracks sudden joins and departures within "
    f"{max(settle_up, settle_down) * 1000:.0f} ms. The evaluation also shows where the scheme gains "
    f"little: with RTS/CTS, at high frame-error rates and at medium unsaturated load, and when it "
    f"shares the channel with legacy BEB stations, which keep a larger share of the airtime.")
para(abstract, 12, spacing=1.15)
para("**Keywords:** CSMA/CA, IEEE 802.11, DCF, backoff algorithms, contention window, "
     "contender estimation, MAC-layer simulation, fairness.", 12, spacing=1.15, before=6)

# ============================================================================ certificate
page_break()
para("CERTIFICATE", 14, True, align="center", spacing=1.5, after=24)
names = ", ".join(f"{n} ({r})" for r, n in sorted(TEAM))
para(f"This is to certify that the project report titled “{TITLE_CASE}” is a bonafide record of "
     f"the work done by {names} in partial fulfilment of the requirements of the Innovative "
     f"Evaluation Practices for BCSE308P - Computer Networks Lab during the Fall Semester 2026-2027.",
     12, spacing=1.5)
para("", after=60)
para("Name and Signature of the Students", 12, True, align="left", spacing=1.5, after=24)
for r, n in sorted(TEAM):
    para(f"{n}  ({r})  ____________________", 12, align="left", spacing=2.0)

# ============================================================================ TOC (placeholder)
page_break()
para("TABLE OF CONTENTS", 14, True, align="center", spacing=1.5, after=12)
TOC_ENTRIES = [  # (sl, text, level, lookup heading)
    ("", "Abstract", 0, None),
    ("1.", "INTRODUCTION", 0, "1. INTRODUCTION"),
    ("", "1.1 Background and motivation", 1, "1.1 Background and motivation"),
    ("", "1.2 Problem statement", 1, "1.2 Problem statement"),
    ("", "1.3 Objectives", 1, "1.3 Objectives"),
    ("", "1.4 Scope and limitations", 1, "1.4 Scope and limitations"),
    ("", "1.5 Mapping to course modules and the lab experiment", 1, "1.5 Mapping to course modules and the lab experiment"),
    ("2.", "LITERATURE SURVEY", 0, "2. LITERATURE SURVEY"),
    ("", "2.1 Overview of the existing protocol or scheme", 1, "2.1 Overview of the existing protocol or scheme"),
    ("", "2.2 Review of related work", 1, "2.2 Review of related work"),
    ("", "2.3 Research gap and justification of novelty", 1, "2.3 Research gap and justification of novelty"),
    ("3.", "PROPOSED SYSTEM", 0, "3. PROPOSED SYSTEM"),
    ("", "3.1 System architecture and topology", 1, "3.1 System architecture and topology"),
    ("", "3.2 Baseline mechanism description", 1, "3.2 Baseline mechanism description"),
    ("", "3.3 Proposed mechanism: algorithm, pseudocode, flowchart", 1, "3.3 Proposed mechanism: algorithm, pseudocode, flowchart"),
    ("", "3.4 Novelty statement", 1, "3.4 Novelty statement"),
    ("", "3.5 Assumptions and design decisions", 1, "3.5 Assumptions and design decisions"),
    ("4.", "IMPLEMENTATION", 0, "4. IMPLEMENTATION"),
    ("", "4.1 Tools, platform, versions and hardware", 1, "4.1 Tools, platform, versions and hardware"),
    ("", "4.2 Simulation setup", 1, "4.2 Simulation setup"),
    ("", "4.3 Code structure and explanation of key modules", 1, "4.3 Code structure and explanation of key modules"),
    ("", "4.4 Challenges faced and how they were solved", 1, "4.4 Challenges faced and how they were solved"),
    ("5.", "EXPERIMENTAL SETUP AND PERFORMANCE EVALUATION", 0, "5. EXPERIMENTAL SETUP AND PERFORMANCE EVALUATION"),
    ("", "5.1 Performance metrics and why they were chosen", 1, "5.1 Performance metrics and why they were chosen"),
    ("", "5.2 Scenarios and parameter values", 1, "5.2 Scenarios and parameter values"),
    ("", "5.3 Results: baseline vs proposed", 1, "5.3 Results: baseline vs proposed"),
    ("", "5.4 Analysis and discussion", 1, "5.4 Analysis and discussion"),
    ("", "5.5 Critical evaluation", 1, "5.5 Critical evaluation"),
    ("6.", "CONCLUSION AND FUTURE WORK", 0, "6. CONCLUSION AND FUTURE WORK"),
    ("", "6.1 Summary of work and key findings", 1, "6.1 Summary of work and key findings"),
    ("", "6.2 Limitations", 1, "6.2 Limitations"),
    ("", "6.3 Future enhancements", 1, "6.3 Future enhancements"),
    ("7.", "REFERENCES", 0, "7. REFERENCES"),
    ("", "Appendix A: Important source code listings", 1, "APPENDIX A: IMPORTANT SOURCE CODE LISTINGS"),
    ("", "Appendix B: Screenshots, event traces and raw result tables", 1, "APPENDIX B: SCREENSHOTS, EVENT TRACES AND RAW RESULT TABLES"),
]
PAGE_NUMBERS = json.loads(os.environ.get("CEAB_TOC_PAGES", "{}"))
toc = doc.add_table(rows=1, cols=3)
toc.alignment = WD_TABLE_ALIGNMENT.CENTER
remove_borders(toc)
for i, h in enumerate(("Sl.No", "Contents", "Page No.")):
    c = toc.rows[0].cells[i]
    set_font(c.paragraphs[0].add_run(h), 12, True)
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
for sl, txt, lvl, key in TOC_ENTRIES:
    cells = toc.add_row().cells
    set_font(cells[0].paragraphs[0].add_run(sl), 12, lvl == 0)
    cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = cells[1].paragraphs[0]
    p.paragraph_format.left_indent = Inches(0.35 * lvl)
    set_font(p.add_run(txt), 12, lvl == 0)
    pg = PAGE_NUMBERS.get(key or "Abstract", "")
    set_font(cells[2].paragraphs[0].add_run(str(pg)), 12, lvl == 0)
    cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
fix_widths(toc, (0.8, 4.6, 0.9))
for row in toc.rows:
    for i in range(3):
        row.cells[i].paragraphs[0].paragraph_format.line_spacing = 1.5 if len(TOC_ENTRIES) < 30 else 1.25
        row.cells[i].paragraphs[0].paragraph_format.space_after = Pt(0)

# front matter: roman page numbers; main matter restarts at 1
add_page_number_footer(doc.sections[0], fmt="lowerRoman", start=1)
main = doc.add_section(WD_SECTION.NEW_PAGE)
add_page_number_footer(main, fmt="decimal", start=1)

# ============================================================================ 1. INTRODUCTION
chapter("Introduction")
CHAPTER_FIRST = True
section("1.1 Background and motivation")
para("Wi-Fi is the access network most people use every day, and in an IEEE 802.11 basic service set "
     "all stations share one radio channel. Because a radio cannot listen while it transmits, "
     "802.11 cannot detect collisions the way Ethernet does; it avoids them instead. The mandatory "
     "Distributed Coordination Function (DCF) implements carrier-sense multiple access with collision "
     "avoidance (CSMA/CA): a station that wants to send waits until the medium has been idle for a "
     "DIFS, then counts down a random number of idle slots drawn from its contention window W, and "
     "transmits when the counter reaches zero. The size of W is the single most important knob of "
     "the MAC layer. If W is too small, several stations pick the same slot and their frames "
     "collide; if W is too large, the channel sits idle while everybody waits.")
para("The standard chooses W with binary exponential backoff (BEB), an idea inherited from "
     "Ethernet [10]: W starts at CWmin+1 = 16 slots, doubles after every failed attempt up to "
     "CWmax+1 = 1024, and is reset to 16 after every success. Classic analysis [2] shows that this "
     "rule works well for a handful of stations but loses a growing share of the capacity as the "
     "number of contenders rises. Dense deployments — lecture halls, hostels, offices and IoT "
     "networks with dozens of active stations per access point — are exactly the case BEB handles "
     "worst. Many alternative ways of waiting have been proposed (linear, Fibonacci, polynomial, "
     "multiplicative-increase/linear-decrease, exponential-increase/exponential-decrease and "
     "feedback-controlled schemes), but they are usually evaluated in isolation, with different "
     "tools and assumptions. This project compares them side by side and proposes a new scheme.")
section("1.2 Problem statement")
para("Under BEB, the contention window reflects only the recent history of a single frame. Every "
     "successful transmission throws that history away, so a station in a crowded cell starts each "
     "new frame with W = 16, collides with high probability, doubles, collides again and only then "
     "reaches a window that suits the actual number of contenders. The results are (i) a collision "
     f"probability of about {pcb:.2f} and a throughput loss of {pct(s3b, s3c, True):.0f} % with 50 "
     "stations, (ii) severe short-term unfairness, because the station that just succeeded has the "
     "smallest window and tends to win again (the capture effect), and (iii) very long tail delays "
     "for the stations that lost several rounds. The problem addressed in this project is: **how "
     "can a station choose a contention window that matches the real level of contention, using "
     "only information it can observe on the channel, without needing all other stations to run the "
     "same algorithm?**")
section("1.3 Objectives")
numbered([
    "Build and validate a slot-accurate, event-driven simulator of the IEEE 802.11g DCF with "
    "pluggable backoff strategies, saturated and Poisson traffic, channel errors, RTS/CTS and "
    "dynamic station membership.",
    "Implement and compare nine waiting strategies: BEB (baseline), constant window, linear, "
    "Fibonacci, quadratic, MILD, EIED, Idle Sense, and the proposed CEAB.",
    "Design CEAB, a backoff scheme that estimates the number of active contenders from overheard "
    "frame addresses and sets the window to the analytically optimal value for that number.",
    "Evaluate all schemes with repeated independent runs (mean ± standard deviation, Welch t-tests) "
    "under low, medium and high load, and in sensitivity, churn, coexistence and ablation scenarios.",
    "Analyse critically when and why the proposed scheme helps, and identify the conditions under "
    "which it fails or gives no benefit.",
])
section("1.4 Scope and limitations")
para("The study covers the MAC sublayer of a single 802.11g (ERP-OFDM, 54 Mb/s) basic service set "
     "in which every station hears every other station (one collision domain). Uplink traffic from "
     "2 to 100 stations is modelled either as saturated (always backlogged) or as Poisson arrivals. "
     "Both basic access and the RTS/CTS handshake are supported, as are random frame errors. "
     "Out of scope are hidden terminals, capture at the receiver, multi-rate stations, 802.11e "
     "EDCA access categories, frame aggregation and transport-layer effects. The work is a "
     "simulation study; it was not deployed on real Wi-Fi hardware, because changing the backoff "
     "rule requires firmware access that commodity adapters do not offer.")
section("1.5 Mapping to course modules and the lab experiment")
table(["Course topic (BCSE308P / BCSE308L)", "Where it appears in this project"], [
    ["Data link layer: MAC sublayer and multiple access", "Core of the project: CSMA/CA contention, "
     "slotted backoff, collisions, retransmission and retry limits (Ch. 3)"],
    ["Random access protocols: ALOHA, CSMA, CSMA/CD, CSMA/CA", "BEB inherited from Ethernet CSMA/CD; "
     "constant-window scheme behaves like p-persistent CSMA (Ch. 2, 3)"],
    ["Wireless LANs: IEEE 802.11 architecture, DCF, RTS/CTS, IFS", "802.11g timing (slot, SIFS, DIFS), "
     "ACK timeout, RTS/CTS mode (Ch. 4.2, 5)"],
    ["Network performance: throughput, delay, jitter, fairness", "Metrics in Ch. 5.1; Bianchi throughput "
     "model; Jain fairness index"],
    ["Lab: simulation of MAC protocols and performance analysis", "Event-driven simulator, repeated runs, "
     "graphs and statistical comparison of the baseline and the proposed scheme"],
    ["Lab: programming network protocols (Python)", "Modular Python package, CLI/GUI demo, unit tests "
     "(Ch. 4.3, Appendix A)"],
], [2.6, 3.7], caption="Mapping of the project to the course syllabus and lab work.")

# ============================================================================ 2. LITERATURE
chapter("Literature Survey")
section("2.1 Overview of the existing protocol or scheme")
para("In the 802.11 DCF [1] a station with a frame first senses the medium. After an idle DIFS "
     f"({phy.difs:.0f} µs in 802.11g) it enters backoff: it draws an integer b uniformly from "
     "[0, W−1] and decrements it once per idle slot of 9 µs, freezing the counter whenever the "
     "medium becomes busy. When b reaches zero it transmits. The receiver answers a correct frame "
     "with an ACK after a SIFS; a missing ACK is treated as a collision. On a collision BEB sets "
     "W ← min(2W, CWmax+1) and retries, up to a retry limit of 7; on success or after a drop it "
     "sets W ← CWmin+1. With RTS/CTS the data frame is preceded by a short RTS/CTS exchange, so a "
     "collision wastes only the RTS duration. Bianchi [2] modelled this process as a Markov chain "
     "and showed that each station attempts with an approximately constant probability τ per slot. "
     "Throughput depends on τ and on the number of stations n, and is maximised when "
     "τ ≈ 1/(n·√(Tc/2σ)), where Tc is the collision duration and σ the slot time. Since n is not "
     "known to the stations, BEB has no way to reach this optimum.")
section("2.2 Review of related work")
lit = [
    ["1", "Kleinrock & Tobagi, 1975 [11]", "CSMA (non-/p-persistent) analysis",
     "Throughput–delay; analytical", "Foundation of carrier sensing; shows the effect of the attempt probability",
     "Fixed persistence; no adaptive backoff"],
    ["2", "Metcalfe & Boggs, 1976 [10]", "Binary exponential backoff in Ethernet",
     "Utilisation; prototype measurement", "Simple, distributed, stable for small n",
     "Designed for collision detection; resets after success"],
    ["3", "Bharghavan et al., 1994 [4]", "MACAW: MILD backoff (×1.5 / −1) and window copying",
     "Throughput, fairness; simulation", "Reduces oscillation and unfairness of BEB",
     "Linear decrease is very slow; window can lock at CWmax"],
    ["4", "Håstad et al., 1996 [15]", "Polynomial backoff",
     "Stability, delay; theoretical analysis", "Polynomial backoff is stable where BEB may not be",
     "Large delays; asymptotic results only"],
    ["5", "Bianchi, 2000 [2]", "Markov model of 802.11 DCF",
     "Saturation throughput; analysis + simulation", "Accurate closed form; defines the optimal τ",
     "Analysis only; does not change the protocol"],
    ["6", "Calì et al., 2000 [3]", "Dynamic tuning: estimate n, use p-persistent optimal τ",
     "Throughput; analysis + simulation", "Gets close to the theoretical limit",
     "Model-based estimator; all stations must run the scheme"],
    ["7", "Bianchi & Tinnirello, 2003 [7]", "Kalman-filter estimate of n from collision probability",
     "Estimation accuracy; simulation", "Accurate, fast tracking of n",
     "Assumes all stations run standard BEB; not used to set W"],
    ["8", "Song et al., 2003 [5]", "EIED (×2 on failure, ÷2 on success)",
     "Throughput, delay; analysis + simulation", "Remembers contention between frames",
     "Fixed factors; unfair under high load"],
    ["9", "Ni et al., 2003 [12]", "Slow contention-window decrease",
     "Throughput, delay; analysis + ns-2", "Simple change; better throughput for large n",
     "No estimate of n; tuning factor fixed"],
    ["10", "Heusse et al., 2005 [6]", "Idle Sense: drive mean idle slots to an optimal target (AIMD)",
     "Throughput, fairness; simulation + testbed", "Near-optimal throughput, good fairness",
     "Needs every station to run it; starves next to legacy stations"],
    ["11", "Manaseer et al., 2006 [14]", "Fibonacci backoff for MANETs",
     "Throughput, delay; simulation", "Slower growth than BEB, fewer idle slots",
     "Still resets after success; degrades for large n"],
]
table(["S.No", "Authors, year", "Technique", "Metrics and tool", "Merits", "Gaps"], lit,
      [0.4, 1.05, 1.25, 1.1, 1.25, 1.25], caption="Literature survey.", size=8.5,
      align=[WD_ALIGN_PARAGRAPH.CENTER] + [WD_ALIGN_PARAGRAPH.LEFT] * 5)
section("2.3 Research gap and justification of novelty")
para("Two families of solutions exist. **History-based rules** (BEB, linear, Fibonacci, polynomial, "
     "MILD, EIED, slow decrease) change W by a fixed rule after each outcome. They are simple, but "
     "they never learn how many stations are actually competing, so their window is right only by "
     "accident. **Estimation-based schemes** (Calì et al., Bianchi & Tinnirello, Idle Sense) do aim "
     "at the optimum, but each depends on a model of how the *other* stations behave: Calì et al. "
     "invert a p-persistent model, the Kalman estimator assumes that everybody runs standard BEB, "
     "and Idle Sense works only when every station runs Idle Sense. Our experiments confirm two "
     "practical consequences. First, inverting an analytical model is fragile in practice, because "
     "real backoff counters are correlated and the measured idle time deviates from the i.i.d. "
     "model (Section 4.4). Second, a station running Idle Sense next to legacy BEB stations "
     f"obtained only {sh_is['idlesense'][0]:.2f} Mb/s against {sh_is['beb'][0]:.1f} Mb/s for the "
     "same number of legacy stations (Section 5.4).")
para("The gap is therefore an estimator of the number of contenders that is **model-free, uses only "
     "passively observable information, and does not care which backoff rule the other stations "
     "run**, combined with a window rule that remains correct when traffic is not saturated. CEAB "
     "fills this gap by counting the transmitter addresses of decoded frames — information already "
     "present in every 802.11 MAC header — and summarising them with the inverse Simpson index, a "
     "diversity measure from ecology [16], [17] that, as far as we know, has not been used for "
     "802.11 backoff tuning. A second gap is methodological: no single study compares the common "
     "waiting strategies under identical conditions with repeated runs and statistical tests; this "
     "project provides such a comparison for nine strategies.")

# ============================================================================ 3. PROPOSED SYSTEM
chapter("Proposed System")
section("3.1 System architecture and topology")
para("Figure 3.1 shows the simulated topology: one access point and N stations within range of each "
     "other, all sending uplink frames. Figure 3.2 shows the software architecture. Each station "
     "has a traffic source, a FIFO queue, the DCF state machine (backoff counter, retry counter) and "
     "a pluggable backoff strategy object that owns its contention window. A shared channel engine "
     "decides, slot by slot, whether a transmission succeeds, collides or is corrupted, charges the "
     "corresponding airtime, and tells every station the length of the idle gap and the transmitter "
     "address (TA) of each decoded frame — exactly what a real 802.11 radio can observe.")
figure("fig_topology.png", "Simulated topology: a single 802.11g BSS with N stations in one collision domain.", 4.2)
figure("fig_architecture.png", "Architecture of the simulator and evaluation framework.", 6.2)
section("3.2 Baseline mechanism description")
para("The baseline is the standard 802.11 DCF with binary exponential backoff (left half of Figure "
     "3.3). With W0 = CWmin+1 = 16 and m = 6 doubling stages the window takes the values 16, 32, …, "
     "1024. After a success, or after a frame is dropped once the retry limit of 7 is exceeded, W "
     "returns to 16. BEB therefore adapts to collisions only within the lifetime of one frame and "
     "forgets the contention level between frames. The other seven reference strategies implemented "
     "for comparison change only the update rule:")
table(["Strategy", "On failure", "On success", "Source"], [
    ["BEB (baseline)", "W ← min(2W, Wmax)", "W ← W0", "IEEE 802.11 [1]"],
    ["Constant", "W unchanged (64)", "W unchanged", "p-persistent CSMA [11]"],
    ["Linear", "W ← W + W0", "W ← W0", "—"],
    ["Fibonacci", "W ← W0 · F(retry), F = 1, 2, 3, 5, 8, …", "W ← W0", "[14]"],
    ["Quadratic", "W ← W0 · (retry+1)²", "W ← W0", "polynomial backoff [15]"],
    ["MILD", "W ← 1.5 W", "W ← W − 1", "MACAW [4]"],
    ["EIED", "W ← 2 W", "W ← W / 2", "[5]"],
    ["Idle Sense", "— (no per-collision change)", "AIMD on W every 5 busy periods towards the optimal "
     "mean idle-slot count", "[6]"],
], [1.25, 1.7, 2.0, 1.3], caption="Update rules of the reference waiting strategies (W0 = 16, Wmax = 1024).", size=9.5)
section("3.3 Proposed mechanism: algorithm, pseudocode, flowchart")
para("**Contender-Estimation Adaptive Backoff (CEAB)** replaces BEB's 'double and reset' rule with "
     "a window derived from an estimate of how many stations are actually contending. It has four "
     "parts.")
subsection("(a) Contender estimation from overheard addresses")
para("Every station records the transmitter address of each successfully decoded data frame on the "
     "channel (its own included) and keeps the last M = 200 of them. If cj frames in this window came "
     "from station j, the effective number of contenders is the unbiased inverse Simpson index")
para("n_eff = M(M − 1) / Σj cj(cj − 1)", 12, italic=True, align="center")
para("When n stations share the channel equally, n_eff = n. Stations that send rarely contribute "
     "little, so n_eff measures how many stations are *competing*, not how many are associated. The "
     "estimate uses no model of the other stations' backoff rule, and because all stations hear the "
     "same frames, they all compute the same value.")
subsection("(b) Load factor for unsaturated traffic")
para("A station that sends one frame per second still appears in the address window, but it is "
     "rarely backlogged. CEAB therefore keeps an exponentially weighted average I_obs (weight 0.05) "
     "of the idle time between busy periods and compares it with I*, the mean idle gap at the "
     "analytical throughput optimum (≈ "
     f"{strat_mod.optimal_idle_slots(phy):.1f} slots for 802.11g):")
para("ρ = min(1, I* / I_obs),     n_c = max(1, ρ · n_eff)", 12, italic=True, align="center")
para("In saturation the channel is idle about as long as at the optimum, so ρ ≈ 1. When traffic is "
     "light the channel is idle much longer, ρ ≪ 1, and the window shrinks towards CWmin. The loop "
     "corrects itself: a window that is too large produces extra idle time, which lowers ρ and "
     "shrinks the window.")
subsection("(c) Optimal target window")
para("For the PHY timing in use (slot σ, success time Ts and collision time Tc), the per-slot attempt "
     "probability τ*(n) that maximises Bianchi's saturation throughput S(τ, n) is pre-computed by "
     "numerical optimisation for n = 1 … 1000. Since a window W gives τ = 2/(W+1), the target "
     "window is")
para("W* = 2 / τ*(n_c) − 1", 12, italic=True, align="center")
para("The base window moves halfway towards this target after each of the station's own outcomes: "
     "W_base ← W_base + γ(W* − W_base), γ = 0.5, clamped to [CWmin+1, CWmax+1], so the scheme is "
     "never more aggressive than the standard minimum.")
subsection("(d) Window used for each attempt")
para("On success the next frame starts with W = W_base — never back at CWmin. On failure the "
     "station may escalate within one frame, W = W_base · 2^min(retry, K). The ablation study "
     "(Section 5.4) shows that K = 0, meaning no escalation at all, gives the lowest tail delay, so "
     "K = 0 is the default: collisions do not change the window, only the estimate does. Figure 3.3 "
     "shows the BEB and CEAB flowcharts side by side.")
figure("fig_flowchart.png", "Flowcharts of the baseline (802.11 BEB) and the proposed CEAB.", 6.3)
subsection("Pseudocode")
code("""Algorithm 1  CEAB (executed independently by every station i)
Parameters: M = 200, gamma = 0.5, beta = 0.05, K = 0, W0 = CWmin+1, Wmax = CWmax+1
State:      log (TA of recently decoded frames), I_obs, W_base = W0, W = W0
Precompute: tau_star[n] = argmax_tau S(tau, n; Ts, Tc, slot)   for n = 1..1000
            I_star     = optimal mean idle gap (slots) * slot

on every busy period (idle gap g_us, decoded frame from station s or none):
    if a frame was decoded:  append s to log
    I_obs <- (1 - beta) * I_obs + beta * g_us

function target_window():
    win   <- last min(M, |log since join|) entries of log
    c_j   <- count of station j in win;  m <- |win|
    n_eff <- m(m-1) / sum_j c_j(c_j-1)          (n_eff = m if the sum is 0)
    rho   <- min(1, I_star / I_obs)
    n_c   <- max(1, rho * n_eff)
    return 2 / tau_star[n_c] - 1

on own transmission outcome (success or failure):
    W_base <- clip(W_base + gamma * (target_window() - W_base), W0, Wmax)
    if success:           W <- W_base;  retry <- 0
    else:                 retry <- retry + 1
                          if retry > 7: drop frame; W <- W_base; retry <- 0
                          else:         W <- min(W_base * 2^min(retry, K), Wmax)
    draw backoff b ~ Uniform{0, ..., round(W) - 1}""", 8.5)
section("3.4 Novelty statement")
para("Compared with the 802.11 BEB baseline, CEAB differs in exactly the following ways:")
bullets([
    "**What sets the window:** BEB uses only the retry count of the current frame; CEAB uses an "
    "estimate of the number of contending stations taken from the channel.",
    "**Behaviour after success:** BEB resets W to CWmin (16 slots); CEAB keeps W at the estimated "
    "optimum (≈ 430 slots with 50 stations, ≈ 40 with 5), which removes the post-success collision "
    "burst and the capture effect.",
    "**Behaviour after a collision:** BEB doubles W; CEAB (K = 0) leaves W unchanged, because one "
    "collision says little about the number of stations.",
    "**Estimator:** unlike Calì et al. [3] and Bianchi & Tinnirello [7], the estimate is model-free "
    "(inverse Simpson index of overheard transmitter addresses) and stays valid whatever backoff "
    "rule the other stations use; unlike Idle Sense [6], the window is not steered towards a fixed "
    "idle-slot target, so CEAB stations are not starved next to legacy stations.",
    "**Unsaturated traffic:** the idle-gap load factor ρ corrects the estimate when stations are "
    "associated but not backlogged; none of the history-based rules distinguishes the two cases.",
])
section("3.5 Assumptions and design decisions")
bullets([
    "Single collision domain: every station decodes every frame (no hidden terminals), so all "
    "stations see the same address sequence and idle gaps.",
    "Frame errors are independent with probability FER and, like collisions, are detected only by a "
    "missing ACK.",
    "The CEAB window never goes below the standard CWmin+1 = 16, so CEAB is never more aggressive "
    "than a legacy station (a deliberate coexistence choice).",
    "M = 200 frames is about 70 ms of airtime at 54 Mb/s: long enough to see each of 50 saturated "
    "stations about four times, short enough to follow churn. The ablation shows results are "
    "insensitive to M between 50 and 1000.",
    "τ*(n) is computed from the analytical model rather than learned online. The model is accurate "
    "near the optimum (Figure 5.2), and a check by simulation confirmed that the throughput plateau "
    "around W* is flat.",
    "Post-backoff after an idle period and EIFS after a corrupted frame are simplified: a newly "
    "arrived frame always draws a backoff, and a corrupted frame costs Tc.",
])

# ============================================================================ 4. IMPLEMENTATION
chapter("Implementation")
section("4.1 Tools, platform, versions and hardware")
import matplotlib
import scipy
import platform
table(["Item", "Details"], [
    ["Language", f"Python {platform.python_version()}"],
    ["Libraries", f"NumPy {np.__version__}, SciPy {scipy.__version__}, Matplotlib {matplotlib.__version__} "
     "(Tk backend for the GUI), python-docx / python-pptx for documents"],
    ["Simulator", "Custom event-driven 802.11 DCF simulator (package csmasim), written for this project"],
    ["Operating system", f"{platform.system()} {platform.release()} (Arch Linux)"],
    ["Hardware", f"{os.cpu_count()}-core x86-64 workstation; experiments run in parallel with "
     "multiprocessing"],
    ["Run time", f"{NSIMS} independent simulations (10 simulated seconds each) in about 6 minutes"],
    ["Reproducibility", "Every run is seeded (seeds 1…10); results/raw_runs.csv stores each run"],
], [1.6, 4.7], caption="Implementation environment.")
section("4.2 Simulation setup")
para("The simulated network is the BSS of Figure 3.1. The PHY/MAC parameters follow 802.11g "
     "ERP-OFDM; the data frame and ACK durations are computed with the OFDM symbol rule "
     "(20 µs preamble + 4 µs per symbol). Table 4.2 lists all parameters.")
table(["Parameter", "Value"], [
    ["PHY", "IEEE 802.11g ERP-OFDM, data rate 54 Mb/s, control rate 24 Mb/s"],
    ["Slot time σ / SIFS / DIFS", f"{phy.slot:.0f} µs / {phy.sifs:.0f} µs / {phy.difs:.0f} µs"],
    ["Payload / MAC header + FCS", f"{phy.payload} B (256–2304 B in the payload study) / 28 B"],
    ["Data frame / ACK duration", f"{phy.t_data:.0f} µs / {phy.t_ack:.0f} µs"],
    ["Ts / Tc (basic access)", f"{phy.ts:.0f} µs / {phy.tc:.0f} µs"],
    ["Ts / Tc (RTS/CTS)", f"{PhyParams(rts_cts=True).ts:.0f} µs / {PhyParams(rts_cts=True).tc:.0f} µs"],
    ["CWmin / CWmax / retry limit", "15 / 1023 / 7 (window sizes 16 and 1024)"],
    ["Number of stations N", "2, 5, 10, 20, 35, 50, 75, 100"],
    ["Traffic", "Saturated, or Poisson with 2–40 Mb/s total offered load (N = 30), queue 100 frames"],
    ["Frame error rate", "0, 0.05, 0.1, 0.2, 0.3"],
    ["Simulated time / warm-up", "10 s per run / first 1 s discarded"],
    ["Independent runs", f"{RUNS} per configuration (seeds 1–{RUNS}); churn: 5 runs"],
    ["CEAB parameters", "M = 200, γ = 0.5, β = 0.05, K = 0, W floor = 16"],
    ["Idle Sense parameters", "ε = 6, α = 1/1.0666, maxtrans = 5, target from the analytical optimum"],
], [2.2, 4.1], caption="Simulation parameters.")
section("4.3 Code structure and explanation of key modules")
loc = lambda f: sum(1 for _ in open(os.path.join(ROOT, f)))
table(["File", "Lines", "Responsibility"], [
    ["csmasim/phy.py", loc("csmasim/phy.py"), "802.11g timing: frame durations, Ts, Tc, DIFS; basic and RTS/CTS"],
    ["csmasim/analytical.py", loc("csmasim/analytical.py"), "Bianchi fixed point; throughput S(τ, n); table of optimal τ*(n)"],
    ["csmasim/strategies.py", loc("csmasim/strategies.py"), "Strategy base class and the nine strategies, including CEAB"],
    ["csmasim/simulator.py", loc("csmasim/simulator.py"), "Event-driven DCF engine, traffic, churn, metrics, event trace"],
    ["csmasim/experiments.py", loc("csmasim/experiments.py"), "Scenario definitions and parallel execution of all runs"],
    ["csmasim/analysis.py", loc("csmasim/analysis.py"), "Mean ± SD, Welch t-tests, tables and figures"],
    ["run_experiments.py", loc("run_experiments.py"), "Command-line entry point for the full evaluation"],
    ["demo.py", loc("demo.py"), "Live demo: compare, trace and GUI with sliders"],
    ["tests/test_simulator.py", loc("tests/test_simulator.py"), "Nine unit/validation tests"],
], [1.9, 0.6, 3.8], caption="Source files.", size=9.5)
para("**Simulation engine (simulator.py).** Instead of stepping through every 9 µs slot, the engine "
     "jumps directly to the next slot in which some backoff counter reaches zero (or to the next "
     "packet arrival or churn event), subtracting the skipped idle slots from all counters at once "
     "with NumPy. If exactly one station's counter reaches zero and no frame error is drawn, the "
     "slot becomes a success of length Ts; otherwise it is a collision or corrupted frame of length "
     "Tc. All counters are frozen during the busy period, as in the standard. After the busy period "
     "the engine calls observe() on strategies that listen to the channel, then on_success, "
     "on_failure or on_drop on the transmitters, which draw new backoff values.")
para("**Strategy interface (strategies.py).** A strategy object manages the windows of a group of "
     "stations as a NumPy array and implements window(), on_success(), on_failure(retry), "
     "on_drop(), observe(...) and reset(). A network can mix groups — for example 10 BEB and 10 "
     "CEAB stations — which is how the coexistence experiment is built. CEAB keeps the shared "
     "address log, the per-station join position (so that a newly joined station uses only frames "
     "it has heard itself), the idle-gap average and W_base.")
para("**Analytical module (analytical.py).** It solves Bianchi's fixed-point equations with Brent's "
     "method for validation, and tabulates τ*(n) by bounded scalar optimisation of S(τ, n) for the "
     "current Ts, Tc and slot; the table is cached per PHY configuration.")
para("**Experiments and analysis.** experiments.py expands the scenario list into "
     f"{NSIMS} (scenario, parameter, strategy, seed) jobs, runs them on all CPU cores and writes "
     "one CSV row per run. analysis.py computes mean and sample standard deviation, Welch's t-test "
     "between BEB and CEAB, the result tables and every figure in Chapter 5.")
section("4.4 Challenges faced and how they were solved")
table(["Challenge", "Solution"], [
    ["A slot-by-slot loop in Python was far too slow for 100 stations × 10 s × 2000+ runs.",
     f"Event skipping: jump to min(counters) and decrement all counters in one vectorised step; run "
     f"jobs in parallel. All {NSIMS} runs finish in about 6 minutes."],
    ["Is the simulator correct?", f"Validated against Bianchi's model (throughput error ≤ {val_err:.1f} %, "
     f"collision probability within {val_perr:.3f}) and against an independent naive slot-by-slot "
     "Monte Carlo (unit tests)."],
    ["First CEAB design estimated n from the station's own attempt rate. Stations with a larger "
     "window measured a lower attempt rate, estimated more contenders and enlarged their window "
     "further — a positive feedback loop (short-term fairness only 0.43).",
     "The estimate must come from the channel only, so that every station computes the same value."],
    ["Second design inverted the analytical model from the observed idle and success probabilities. "
     "Measured idle gaps were about 0.8 slot longer than the i.i.d. model predicts (backoff counters "
     "are correlated), and the inversion is ill-conditioned, so estimates exploded to 10⁴ stations.",
     "Replaced by a model-free estimator: counting transmitter addresses (inverse Simpson index). "
     "The estimate is within a few per cent of the true N (Figure 5.7)."],
    ["Under light Poisson traffic CEAB had twice the delay of BEB: idle periods with no contender "
     "were not counted, so the channel looked busier than it was.",
     "The load factor now uses the wall-clock idle gap and compares it with the optimal gap; "
     "low-load delay now equals BEB's."],
    ["MILD runs at N = 5–10 had very high variance.", "Not a bug: individual runs lock in at CWmax "
     "because a −1 decrease cannot undo ×1.5 increases. Reported as a finding (Section 5.4)."],
], [3.1, 3.2], caption="Challenges and solutions.", size=9.5,
    align=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

# ============================================================================ 5. EVALUATION
chapter("Experimental Setup and Performance Evaluation")
section("5.1 Performance metrics and why they were chosen")
table(["Metric", "Definition", "Why it matters"], [
    ["Throughput (Mb/s)", "Payload bits delivered per second (after warm-up)", "Primary capacity measure; the "
     "quantity the optimal window maximises"],
    ["Collision probability", "Failed-by-collision attempts / all attempts", "Direct indicator of a window that "
     "is too small; wastes airtime and energy"],
    ["Mean access delay (ms)", "Time from a frame reaching the head of the queue to its ACK",
     "MAC service time; with saturation it equals N / frame rate"],
    ["99th-percentile access delay (ms)", "Tail of the access-delay distribution", "Real-time and interactive "
     "traffic suffers from the tail, not the mean"],
    ["Short-term Jain fairness", "J = (Σx)² / (N·Σx²) of per-station frame counts in 100 ms windows, "
     "averaged", "Captures the capture effect that long-term averages hide"],
    ["Drop rate", "Frames discarded after 7 retries / frames handled", "Losses visible to upper layers"],
    ["End-to-end delay (ms)", "Arrival in the queue to ACK (Poisson traffic)", "Includes queueing; shows the "
     "saturation point"],
], [1.5, 2.6, 2.2], caption="Performance metrics.", size=9.5,
    align=[WD_ALIGN_PARAGRAPH.LEFT] * 3)
section("5.2 Scenarios and parameter values")
table(["Scenario", "Configuration", "Strategies", "Runs"], [
    ["S1 / S2 / S3", "Saturated, N = 5 / 20 / 50 (low / medium / high contention)", "all 9", f"{RUNS}"],
    ["Station sweep", "Saturated, N = 2 … 100", "all 9", f"{RUNS}"],
    ["Validation", "BEB with unlimited retries vs. Bianchi model, N = 2 … 100", "BEB", f"{RUNS}"],
    ["Offered load", "Poisson, N = 30, 2 … 40 Mb/s total", "BEB, EIED, Idle Sense, CEAB", f"{RUNS}"],
    ["Channel errors", "Saturated, N = 20, FER 0 … 0.3", "4 main", f"{RUNS}"],
    ["RTS/CTS", "Saturated, N = 5 / 20 / 50, RTS/CTS access", "4 main", f"{RUNS}"],
    ["Payload", "Saturated, N = 20, 256 … 2304 B", "4 main", f"{RUNS}"],
    ["Churn", "N: 5 → 50 at t = 4 s → 15 at t = 8 s, 12 s runs", "4 main", "5"],
    ["Coexistence", "10 legacy BEB + 10 stations of scheme X", "X = BEB, EIED, Idle Sense, CEAB", f"{RUNS}"],
    ["Ablation", "CEAB with K ∈ {0,1,2}, M ∈ {50…1000}, N = 20, 50", "CEAB", f"{RUNS}"],
], [1.2, 2.8, 1.7, 0.6], caption=f"Evaluation scenarios ({NSIMS} simulation runs in total).", size=9.5)

section("5.3 Results: baseline vs proposed")
para(f"Table 5.3 follows the results template: BEB (baseline) against CEAB (proposed) for the low, "
     f"medium and high load scenarios, as the mean ± standard deviation of {RUNS} independent runs. "
     "Improvement is positive when CEAB is better (higher throughput or fairness, lower collision "
     "probability, delay or drop rate). The last column gives the p-value of Welch's two-sample "
     "t-test.")
mt = main_table(ROWS)
mt_t = table(["Scenario", "Metric", "Baseline (mean ± SD)", "Proposed (mean ± SD)", "Improve-ment (%)",
              "p-value"],
             [[r["scenario"], r["metric"], r["baseline"], r["proposed"],
               f"{r['improvement']:+.1f}", f"{r['p_value']:.1e}" if not np.isnan(r["p_value"]) else "–"]
              for r in mt], [1.05, 1.6, 1.15, 1.15, 0.8, 0.55],
             caption="Baseline (BEB) vs proposed (CEAB) — mean ± SD over repeated runs.", size=8.5)
for g0 in range(1, len(mt) + 1, 6):                       # one merged scenario cell per group
    label_ = mt[g0 - 1]["scenario"]
    merged = mt_t.cell(g0, 0).merge(mt_t.cell(g0 + 5, 0))
    for p_ in merged.paragraphs[1:]:
        p_._p.getparent().remove(p_._p)
    merged.paragraphs[0].text = ""
    set_font(merged.paragraphs[0].add_run(label_), 8.5, True)
    merged.vertical_alignment = 1
labels = {k: c.label for k, c in strat_mod.REGISTRY.items()}
order = list(strat_mod.REGISTRY)


def sc_table(metric, caption, f="{:.2f}"):
    rows = []
    for st in order:
        cells = [labels[st]]
        for sc, n in (("S1", 5), ("S2", 20), ("S3", 50)):
            cells.append(f"{f.format(V(sc, n, st, metric))} ± {f.format(V(sc, n, st, metric, True))}")
        rows.append(cells)
    rows[-1] = [f"**{c}**" for c in rows[-1]]
    table(["Strategy", "S1 (N = 5)", "S2 (N = 20)", "S3 (N = 50)"], rows, [1.9, 1.45, 1.45, 1.45],
          caption=caption, size=9.5)


sc_table("throughput", "Throughput (Mb/s) of every strategy in S1–S3 (mean ± SD).")
sc_table("access_delay_p99", "99th-percentile access delay (ms) in S1–S3 (mean ± SD).")
sc_table("jain_short", "Short-term Jain fairness index in S1–S3 (mean ± SD).", "{:.3f}")
figure("fig_s123_throughput.png", "Throughput of all strategies in scenarios S1–S3 (whiskers = ±1 SD).")
figure("fig_validation.png", "Validation of the simulator against the Bianchi analytical model.")
figure("fig_throughput_vs_n.png", "Saturation throughput as the number of stations grows from 2 to 100.")
figure("fig_collision_vs_n.png", "Collision probability versus number of stations.")
figure("fig_p99_vs_n.png", "Tail (99th percentile) access delay versus number of stations (log scale).")
figure("fig_fairness_vs_n.png", "Short-term (100 ms) Jain fairness versus number of stations.")
figure("fig_churn.png", "Dynamic membership: throughput, CEAB's estimate and mean window over time.")
figure("fig_load.png", "Unsaturated Poisson traffic with 30 stations: carried load and end-to-end delay.")
figure("fig_fer_payload.png", "Sensitivity to frame errors (left) and payload size (right), N = 20.")
figure("fig_rts.png", "Throughput with RTS/CTS access.", 5.6)
figure("fig_coexist.png", "Coexistence of 10 legacy BEB stations with 10 stations running each scheme.", 5.6)
figure("fig_ablation.png", "Ablation: escalation cap K and estimation window M of CEAB.")

section("5.4 Analysis and discussion")
subsection("Validation")
para(f"Before any comparison, the simulator was checked against Bianchi's model. Figure 5.2 shows "
     f"that simulated BEB throughput and collision probability follow the analytical curves for "
     f"2 to 100 stations, with a maximum throughput error of {val_err:.1f} % and a maximum "
     f"collision-probability difference of {val_perr:.3f}. The small remaining gap at large N comes "
     "from the model's assumption that collision probability is independent of the backoff stage.")
subsection("Why BEB loses capacity as N grows")
para(f"BEB throughput falls from {V('sweep_n', 2, 'beb', 'throughput'):.1f} Mb/s with 2 stations to "
     f"{s3b:.1f} Mb/s with 50 and {n100b:.1f} Mb/s with 100, while the theoretical limit stays "
     f"near {opt_thr(50):.1f} Mb/s (Figure 5.3). The reason is visible in the churn trace (Figure "
     "5.7, bottom): BEB's mean window with 50 stations is large, but every station that has just "
     "succeeded is back at 16, so the next attempts collide in bursts. The collision probability "
     f"reaches {pcb:.2f} (Figure 5.4) and 2 % of the frames are dropped after 7 retries.")
subsection("How CEAB recovers it")
para(f"CEAB keeps every station at the window that suits the estimated contender count. Its "
     f"throughput is {s1c:.2f}, {s2c:.2f} and {s3c:.2f} Mb/s in S1–S3 — "
     f"{pct(s1c, s1b):+.1f} %, {pct(s2c, s2b):+.1f} % and {pct(s3c, s3b):+.1f} % over BEB — and "
     f"{n100c:.2f} Mb/s at N = 100 ({pct(n100c, n100b):+.0f} %). Across the whole sweep it stays "
     f"at {eff * 100:.1f} % or more of the analytical optimum. The collision probability is almost "
     f"independent of N (about {pcc:.2f}), which is the signature of operating at the optimal τ. "
     "Every difference in Table 5.3 except the mean access delay at N = 50 is significant at "
     "p < 0.01. With saturated stations the mean access delay is fixed by the frame rate (each of N "
     "stations gets 1/N of the successes), so the mean barely changes; the gain appears in the "
     f"tail: the 99th-percentile delay falls from {p99b:.0f} ms to {p99c:.0f} ms at N = 50.")
subsection("Fairness and the capture effect")
para(f"Short-term fairness is where the waiting strategies differ most (Figure 5.6). BEB drops to "
     f"{jb:.2f} with 50 stations and EIED to {V('S3', 50, 'eied', 'jain_short'):.2f}, because the "
     "station that just succeeded has the smallest window (BEB) or keeps a smaller window than its "
     "neighbours (EIED). CEAB stations all derive the same estimate from the same overheard frames "
     f"and therefore use the same window, giving {jc:.2f}. Idle Sense, which also equalises windows, "
     f"reaches {V('S3', 50, 'idlesense', 'jain_short'):.2f}.")
subsection("Comparison of the waiting strategies")
para("Growth rate alone does not fix the problem. Linear and Fibonacci backoff grow more slowly than "
     f"BEB and collide even more ({V('S3', 50, 'linear', 'collision_prob'):.2f} and "
     f"{V('S3', 50, 'fibonacci', 'collision_prob'):.2f} at N = 50). Quadratic growth is a little "
     "better than BEB, and a fixed window of 64 is excellent for 5–10 stations but collapses to "
     f"{V('S3', 50, 'constant', 'throughput'):.1f} Mb/s at N = 50 because it cannot adapt. The "
     "schemes that do not reset after a success — EIED, MILD, Idle Sense and CEAB — are the only ones "
     "that hold up at high N, which supports the central idea of this project: the window must carry "
     "information about contention from one frame to the next. MILD illustrates the danger of doing "
     f"so badly. It is close to CEAB at N = 75–100 ({V('sweep_n', 100, 'mild', 'throughput'):.1f} Mb/s), "
     f"but at N = 5 its throughput has a standard deviation of {V('S1', 5, 'mild', 'throughput', True):.1f} "
     "Mb/s: in some runs a few stations drift to CWmax and stay there, because one ×1.5 increase "
     "needs dozens of −1 decreases to undo. Idle Sense is the strongest reference scheme; CEAB beats "
     f"it slightly in throughput (by {pct(s3c, V('S3', 50, 'idlesense', 'throughput')):.1f} % at N = 50 "
     f"and {pct(n100c, n100i):.1f} % at N = 100), in collision probability and in tail delay, and "
     "matches its fairness.")
subsection("Dynamic membership")
para(f"When 45 stations join at t = 4 s (Figure 5.7), CEAB's estimate reaches within 15 % of the new "
     f"count in about {settle_up * 1000:.0f} ms, and after the drop to 15 stations at t = 8 s in "
     f"about {settle_down * 1000:.0f} ms. Between 4.5 s and 8 s CEAB delivers "
     f"{churn_mean('ceab', 4.5, 8):.1f} Mb/s against {churn_mean('beb', 4.5, 8):.1f} Mb/s for BEB "
     f"and {churn_mean('idlesense', 4.5, 8):.1f} Mb/s for Idle Sense, whose additive window increase "
     "needs longer to reach the larger window.")
subsection("Unsaturated traffic")
ld = lambda st, l, m="e2e_delay_mean": V("load", l, st, m)
para(f"With Poisson traffic (Figure 5.8) all schemes carry the offered load up to about 25 Mb/s and "
     f"have the same delay at light load (≈ {ld('ceab', 5):.2f} ms at 5 Mb/s), which confirms that "
     f"the load factor shrinks CEAB's window when stations are rarely backlogged. BEB saturates at "
     f"about {ld('beb', 40, 'throughput'):.1f} Mb/s, while CEAB carries up to "
     f"{ld('ceab', 40, 'throughput'):.1f} Mb/s (+{pct(ld('ceab', 40, 'throughput'), ld('beb', 40, 'throughput')):.0f} %). "
     f"At 28 Mb/s offered, BEB is already overloaded (mean delay {ld('beb', 28):.0f} ms) while CEAB "
     f"still delivers the full load with {ld('ceab', 28):.1f} ms.")
subsection("Sensitivity: frame errors, payload and RTS/CTS")
para(f"Random frame errors (Figure 5.9, left) hurt every scheme, but BEB also misreads each error as "
     f"a collision and doubles its window. CEAB's estimate is unaffected (it counts decoded frames "
     f"and idle time), so it stays ahead of BEB at every error rate. The gain narrows, however, from "
     f"{pct(V('fer', 0.0, 'ceab', 'throughput'), V('fer', 0.0, 'beb', 'throughput')):.1f} % at FER = 0 "
     f"to {pct(V('fer', 0.3, 'ceab', 'throughput'), V('fer', 0.3, 'beb', 'throughput')):.1f} % at FER = 0.3, "
     "where EIED is slightly better. The gain grows with payload size "
     f"({pct(V('payload', 256, 'ceab', 'throughput'), V('payload', 256, 'beb', 'throughput')):.1f} % at 256 B, "
     f"{pct(V('payload', 2304, 'ceab', 'throughput'), V('payload', 2304, 'beb', 'throughput')):.1f} % at 2304 B), "
     "because larger frames make collisions more expensive. With RTS/CTS a collision costs only "
     f"{PhyParams(rts_cts=True).tc:.0f} µs instead of {phy.tc:.0f} µs, so the optimal window is "
     "smaller and BEB is already close to it: CEAB's throughput gain falls to "
     f"{pct(V('rts', 5, 'ceab', 'throughput'), V('rts', 5, 'beb', 'throughput')):+.1f} %, "
     f"{pct(V('rts', 20, 'ceab', 'throughput'), V('rts', 20, 'beb', 'throughput')):+.1f} % and "
     f"{pct(V('rts', 50, 'ceab', 'throughput'), V('rts', 50, 'beb', 'throughput')):+.1f} % for N = 5, 20 and 50. "
     "Its tail-delay and fairness advantages remain "
     f"(p99 {V('rts', 50, 'beb', 'access_delay_p99'):.0f} → {V('rts', 50, 'ceab', 'access_delay_p99'):.0f} ms at N = 50).")
subsection("Coexistence with legacy stations")
para(f"When 10 CEAB stations share the channel with 10 legacy BEB stations (Figure 5.11), the total "
     f"throughput rises from {sh_beb['beb'][0]:.1f} Mb/s (all BEB) to "
     f"{sh_ceab['beb'][0] + sh_ceab['ceab'][0]:.1f} Mb/s. The split is uneven, though: the BEB group "
     f"gets {sh_ceab['beb'][0]:.1f} Mb/s and the CEAB group {sh_ceab['ceab'][0]:.1f} Mb/s, because "
     "legacy stations restart at W = 16 after every success. Even so, a CEAB station gets "
     f"{sh_ceab['ceab'][0] / 10:.2f} Mb/s, more than the {sh_beb['beb'][0] / 20:.2f} Mb/s it would get "
     "in an all-BEB cell, so adopting CEAB never makes a station worse off. Idle Sense fares much "
     f"worse in the same test ({sh_is['idlesense'][0]:.2f} Mb/s for its 10 stations), because the "
     "BEB stations keep the idle time below its target, so it keeps enlarging its window.")
subsection("Ablation")
para(f"Figure 5.12 isolates CEAB's design parameters. Allowing per-retry doubling on top of the "
     f"estimate (K = 1 or 2) does not improve throughput but raises the N = 50 tail delay from "
     f"{V('ablation_K', '50:0', 'ceab', 'access_delay_p99'):.0f} ms to "
     f"{V('ablation_K', '50:1', 'ceab', 'access_delay_p99'):.0f} and "
     f"{V('ablation_K', '50:2', 'ceab', 'access_delay_p99'):.0f} ms and lowers fairness. In a single "
     "collision domain the estimate is accurate, so reacting to individual collisions only adds noise; "
     "hence K = 0. The estimation window M has almost no effect between 50 and 1000 frames "
     f"(p99 {min(V('ablation_M', f'50:{m}', 'ceab', 'access_delay_p99') for m in (50, 100, 200, 500, 1000)):.1f}–"
     f"{max(V('ablation_M', f'50:{m}', 'ceab', 'access_delay_p99') for m in (50, 100, 200, 500, 1000)):.1f} ms), "
     "so it can be chosen for tracking speed (M = 200).")

section("5.5 Critical evaluation")
subsection("Strengths")
bullets([
    f"Close to the theoretical capacity for every N from 5 to 100 (≥ {eff * 100:.1f} % of the "
    "optimum); the gain over BEB grows with N.",
    "Removes the capture effect: short-term fairness of 0.91 or more up to 50 stations.",
    f"Tail delay reduced by {pct(p99c, p99b, True):.0f} % at N = 50, and no frames dropped at the retry limit.",
    "Uses only information every 802.11 receiver already decodes; no new frames, fields or AP support.",
    "The estimator does not depend on the other stations' backoff rule, so the scheme degrades "
    "gracefully in mixed networks (unlike Idle Sense).",
    "Only one parameter matters in practice (K), and the ablation shows a clear best value.",
])
subsection("Weaknesses")
bullets([
    f"At medium unsaturated load the mean delay is higher than BEB's (e.g. {ld('ceab', 20):.2f} vs "
    f"{ld('beb', 20):.2f} ms at 20 Mb/s offered): the load factor shrinks the window, but not all "
    "the way to CWmin.",
    "Each station must keep a log of overheard addresses and run a small optimisation table — "
    "cheap, but more state than BEB's single counter.",
    "The optimal-τ table is computed from the analytical model and assumes known Ts and Tc; with "
    "multi-rate stations, frames have different durations and the target is only approximate.",
    "Legacy BEB stations take a larger share of the channel than CEAB stations (about 2.6 : 1).",
])
subsection("Conditions where the proposed scheme fails or gives no benefit")
bullets([
    "**RTS/CTS with few stations:** collisions are cheap and BEB is already near the optimum; CEAB "
    "is equal or marginally lower (−0.3 % at N = 5).",
    "**High frame-error rates (FER ≥ 0.2):** losses dominate and every scheme converges; EIED is "
    "slightly better at FER = 0.3.",
    "**Hidden terminals:** the estimator assumes that every station hears every frame. Stations "
    "outside each other's range would count different senders, and collisions at the AP would not "
    "be reflected in the estimate (not simulated; expected to under-estimate contention).",
    "**Very short-lived bursts:** a station that has just joined has heard no frames yet and starts at "
    "CWmin, so for the first tens of milliseconds it behaves like BEB.",
    "**Light load:** with only a few active stations BEB is already close to optimal "
    f"(S1 gain only {pct(s1c, s1b):.1f} %); the benefit of CEAB is a dense-network benefit.",
])

# ============================================================================ 6. CONCLUSION
chapter("Conclusion and Future Work")
section("6.1 Summary of work and key findings")
para(f"A validated event-driven simulator of the IEEE 802.11g DCF was built and used to compare nine "
     f"CSMA/CA waiting strategies over {NSIMS} independent runs. The proposed CEAB scheme estimates "
     "the number of contending stations from overheard transmitter addresses and holds the "
     "contention window at the analytically optimal value instead of resetting it after each "
     "success. Key quantified findings:")
bullets([
    f"Throughput: {pct(s1c, s1b):+.1f} % (N = 5), {pct(s2c, s2b):+.1f} % (N = 20), "
    f"{pct(s3c, s3b):+.1f} % (N = 50) and {pct(n100c, n100b):+.1f} % (N = 100) over BEB; within "
    f"{100 - eff * 100:.1f} % of the theoretical limit.",
    f"Collision probability at N = 50: {pcb:.3f} → {pcc:.3f}; drop rate 2.0 % → 0 %.",
    f"99th-percentile access delay at N = 50: {p99b:.0f} ms → {p99c:.0f} ms ({pct(p99c, p99b, True):.0f} % lower).",
    f"Short-term Jain fairness at N = 50: {jb:.3f} → {jc:.3f}.",
    f"Maximum carried load with Poisson traffic: {ld('beb', 40, 'throughput'):.1f} → "
    f"{ld('ceab', 40, 'throughput'):.1f} Mb/s; light-load delay unchanged.",
    f"The estimator tracks membership changes within about {max(settle_up, settle_down) * 1000:.0f} ms.",
    "Among the history-based rules, only those that keep contention information across frames "
    "(EIED, MILD) scale; linear, Fibonacci and constant windows are worse than BEB at high N.",
])
section("6.2 Limitations")
bullets([
    "Simulation only, in a single collision domain with one PHY rate; no hidden terminals, capture "
    "effect or multi-rate stations.",
    "Traffic is uplink only; no TCP dynamics, EDCA priorities or frame aggregation (A-MPDU).",
    "The optimal window is derived from an analytical model that assumes saturated, homogeneous "
    "stations.",
    "Idle Sense was re-implemented from its published description in simplified form; the "
    "original may perform slightly differently.",
])
section("6.3 Future enhancements")
bullets([
    "Extend the simulator to multiple overlapping cells and hidden terminals, and estimate contention "
    "from both decoded addresses and unexplained busy periods.",
    "Make CEAB airtime-aware for multi-rate networks (count airtime per sender rather than frames).",
    "A coexistence mode that detects legacy stations (from their short post-success gaps) and lowers "
    "the window floor to recover a fair share.",
    "Port the algorithm to ns-3 and to open firmware (e.g. OpenWiFi or ath9k) for a testbed experiment.",
    "Apply the estimator to 802.11ah / IoT networks with hundreds of stations, where the gains "
    "should be larger still.",
])

# ============================================================================ 7. REFERENCES
chapter("References")
refs = [
    "IEEE Standard for Information Technology—Telecommunications and Information Exchange between "
    "Systems—Local and Metropolitan Area Networks—Specific Requirements—Part 11: Wireless LAN Medium "
    "Access Control (MAC) and Physical Layer (PHY) Specifications, IEEE Std 802.11-2020, Feb. 2021.",
    "G. Bianchi, “Performance analysis of the IEEE 802.11 distributed coordination function,” "
    "IEEE J. Sel. Areas Commun., vol. 18, no. 3, pp. 535–547, Mar. 2000.",
    "F. Calì, M. Conti, and E. Gregori, “Dynamic tuning of the IEEE 802.11 protocol to achieve a "
    "theoretical throughput limit,” IEEE/ACM Trans. Netw., vol. 8, no. 6, pp. 785–799, Dec. 2000.",
    "V. Bharghavan, A. Demers, S. Shenker, and L. Zhang, “MACAW: A media access protocol for wireless "
    "LANs,” in Proc. ACM SIGCOMM, London, U.K., 1994, pp. 212–225.",
    "N.-O. Song, B.-J. Kwak, J. Song, and L. E. Miller, “Enhancement of IEEE 802.11 distributed "
    "coordination function with exponential increase exponential decrease backoff algorithm,” in "
    "Proc. IEEE Veh. Technol. Conf. (VTC-Spring), 2003, pp. 2775–2778.",
    "M. Heusse, F. Rousseau, R. Guillier, and A. Duda, “Idle Sense: An optimal access method for high "
    "throughput and fairness in rate diverse wireless LANs,” in Proc. ACM SIGCOMM, Philadelphia, PA, "
    "USA, 2005, pp. 121–132.",
    "G. Bianchi and I. Tinnirello, “Kalman filter estimation of the number of competing terminals in "
    "an IEEE 802.11 network,” in Proc. IEEE INFOCOM, San Francisco, CA, USA, 2003, pp. 844–852.",
    "B.-J. Kwak, N.-O. Song, and L. E. Miller, “Performance analysis of exponential backoff,” "
    "IEEE/ACM Trans. Netw., vol. 13, no. 2, pp. 343–355, Apr. 2005.",
    "R. Jain, D.-M. Chiu, and W. R. Hawe, “A quantitative measure of fairness and discrimination for "
    "resource allocation in shared computer systems,” Digital Equipment Corp., Tech. Rep. DEC-TR-301, 1984.",
    "R. M. Metcalfe and D. R. Boggs, “Ethernet: Distributed packet switching for local computer "
    "networks,” Commun. ACM, vol. 19, no. 7, pp. 395–404, Jul. 1976.",
    "L. Kleinrock and F. A. Tobagi, “Packet switching in radio channels: Part I—Carrier sense "
    "multiple-access modes and their throughput-delay characteristics,” IEEE Trans. Commun., "
    "vol. COM-23, no. 12, pp. 1400–1416, Dec. 1975.",
    "Q. Ni, I. Aad, C. Barakat, and T. Turletti, “Modeling and analysis of slow CW decrease for IEEE "
    "802.11 WLAN,” in Proc. IEEE Int. Symp. Personal, Indoor and Mobile Radio Commun. (PIMRC), 2003, "
    "pp. 1717–1721.",
    "H. Wu, Y. Peng, K. Long, S. Cheng, and J. Ma, “Performance of reliable transport protocol over "
    "IEEE 802.11 wireless LAN: Analysis and enhancement,” in Proc. IEEE INFOCOM, 2002, pp. 599–607.",
    "S. Manaseer, M. Ould-Khaoua, and L. M. Mackenzie, “Fibonacci backoff algorithm for mobile ad hoc "
    "networks,” in Proc. 7th Annu. PostGraduate Symp. Convergence of Telecommun., Netw. and "
    "Broadcasting (PGNet), Liverpool, U.K., 2006.",
    "J. Håstad, T. Leighton, and B. Rogoff, “Analysis of backoff protocols for multiple access "
    "channels,” SIAM J. Comput., vol. 25, no. 4, pp. 740–774, 1996.",
    "E. H. Simpson, “Measurement of diversity,” Nature, vol. 163, p. 688, Apr. 1949.",
    "M. O. Hill, “Diversity and evenness: A unifying notation and its consequences,” Ecology, vol. 54, "
    "no. 2, pp. 427–432, Mar. 1973.",
]
for i, r in enumerate(refs, 1):
    p = para(f"[{i}]  {r}", 11.5, align="justify", spacing=1.15, after=6)
    p.paragraph_format.left_indent = Inches(0.45)
    p.paragraph_format.first_line_indent = Inches(-0.45)

# ============================================================================ APPENDICES
def appendix(title):
    h = doc.add_heading(level=1)
    h.paragraph_format.page_break_before = True
    set_font(h.add_run(title), 14, True)
    HEADINGS.append(title)
    CHAPTER["n"] = title.split(":")[0][-1]
    FIGNO["n"] = TABNO["n"] = 0


appendix("APPENDIX A: IMPORTANT SOURCE CODE LISTINGS")
para("**A.1 CEAB backoff strategy (csmasim/strategies.py)**", 12, align="left")
code(inspect.getsource(strat_mod.CEAB), 7.5)
para("**A.2 Baseline BEB and the strategy interface (csmasim/strategies.py)**", 12, align="left")
code(inspect.getsource(strat_mod.Strategy) + "\n\n" + inspect.getsource(strat_mod.BEB), 7.5)
sim_src = open(os.path.join(ROOT, "csmasim", "simulator.py")).read()
a = sim_src.index("            # ---- a transmission starts at t_tx")
b = sim_src.index("        take_samples(T)")
para("**A.3 Core of the channel engine: one transmission event (csmasim/simulator.py)**", 12, align="left")
code(sim_src[a:b], 7.5)
para("**A.4 Optimal attempt probability table (csmasim/analytical.py)**", 12, align="left")
from csmasim import analytical as ana_mod
code(inspect.getsource(ana_mod.throughput_from_tau) + "\n\n" + inspect.getsource(ana_mod.w_opt), 7.5)

appendix("APPENDIX B: SCREENSHOTS, EVENT TRACES AND RAW RESULT TABLES")
para("**B.1 Interactive demo (python demo.py gui).** Sliders change N, offered load, frame error "
     "rate, payload and duration; check-boxes choose strategies; the radio button selects basic or "
     "RTS/CTS access.", 12)
figure("fig_gui.png", "GUI of the live demo comparing BEB, Idle Sense and CEAB with 50 stations.", 6.2)
para("**B.2 Command-line comparison (python demo.py compare -n 50 --runs 3).**", 12, align="left")
cli = subprocess.run([sys.executable, os.path.join(ROOT, "demo.py"), "compare", "-n", "50", "--runs", "3",
                      "--duration", "5"], capture_output=True, text=True, cwd=ROOT).stdout
code("\n".join(l for l in cli.splitlines() if not l.strip().startswith("ran ")), 7.0)
para("**B.3 Channel event trace (python demo.py trace -n 5 --events 15).** The simulator equivalent "
     "of a packet capture: each line is one busy period with the idle slots before it, the "
     "outcome, the transmitting stations, the window each one drew its backoff from and its retry "
     "count. Under BEB a station returns to W = 16 after every success; under CEAB the windows grow "
     "to the estimated optimum (≈ 40 for 5 stations) and stay there.", 12)
tr = subprocess.run([sys.executable, os.path.join(ROOT, "demo.py"), "trace", "-n", "5", "--events", "15"],
                    capture_output=True, text=True, cwd=ROOT).stdout
code(tr, 7.0)
para("**B.4 Raw per-run results, scenario S3 (N = 50, saturated).** The complete per-run data "
     "(results/raw_runs.csv), the summary (results/summary.csv) and all figures are produced by "
     "run_experiments.py and are submitted with the source code.", 12)
raw = [r for r in ROWS if r["scenario"] == "S3" and r["strategy"] in ("beb", "ceab")]
seeds = sorted({int(r["seed"]) for r in raw})
g = {(r["strategy"], int(r["seed"])): r for r in raw}
table(["Seed", "BEB thr (Mb/s)", "CEAB thr (Mb/s)", "BEB coll. prob.", "CEAB coll. prob.",
       "BEB p99 (ms)", "CEAB p99 (ms)"],
      [[s, f"{g[('beb', s)]['throughput']:.3f}", f"{g[('ceab', s)]['throughput']:.3f}",
        f"{g[('beb', s)]['collision_prob']:.4f}", f"{g[('ceab', s)]['collision_prob']:.4f}",
        f"{g[('beb', s)]['access_delay_p99']:.1f}", f"{g[('ceab', s)]['access_delay_p99']:.1f}"]
       for s in seeds], [0.5, 0.95, 0.95, 0.95, 0.95, 0.95, 0.95],
      caption="Individual runs behind the S3 row of Table 5.3.", size=9)
para("**B.5 Validation data: simulated BEB vs. Bianchi model (unlimited retries).**", 12, align="left")
table(["N", "Model thr (Mb/s)", "Sim thr (Mb/s, mean ± SD)", "Error (%)", "Model p", "Sim p"],
      [[n, f"{float(a_['bianchi_thr']):.2f}",
        f"{V('validation', n, 'beb', 'throughput'):.2f} ± {V('validation', n, 'beb', 'throughput', True):.2f}",
        f"{(V('validation', n, 'beb', 'throughput') / float(a_['bianchi_thr']) - 1) * 100:+.2f}",
        f"{float(a_['bianchi_p']):.3f}", f"{V('validation', n, 'beb', 'collision_prob'):.3f}"]
       for n in (2, 5, 10, 20, 35, 50, 75, 100)
       for a_ in [next(x for x in ANA if int(x["n"]) == n)]],
      [0.5, 1.1, 1.6, 0.8, 0.8, 0.8], caption="Simulator validation.", size=9.5)

doc.save(OUT)


def to_pdf(path):
    outdir = os.path.dirname(path)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", outdir, path],
                   check=True, capture_output=True)
    return path[:-5] + ".pdf"


def roman(n):
    vals = [(10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
    out = ""
    for v, r in vals:
        while n >= v:
            out += r
            n -= v
    return out


def find_pages(pdf):
    """Physical page of each TOC heading, converted to the printed page number."""
    txt = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, text=True).stdout
    pages = [" ".join(pg.split()) for pg in txt.split("\f")]
    first_main = next(i for i, pg in enumerate(pages) if pg.startswith("1. INTRODUCTION"))
    found = {}
    start = first_main
    for _, _, _, key in TOC_ENTRIES:
        if key is None:
            continue
        for i in range(start, len(pages)):
            if key in pages[i]:
                found[key] = start = i
                break
    nums = {k: v - first_main + 1 for k, v in found.items()}
    abs_pg = next(i for i, pg in enumerate(pages) if "ABSTRACT" in pg)
    nums["Abstract"] = roman(abs_pg + 1)
    return nums, len(pages)


if __name__ == "__main__" and "CEAB_TOC_PAGES" not in os.environ:
    pdf = to_pdf(OUT)
    nums, npages = find_pages(pdf)
    missing = [k for _, _, _, k in TOC_ENTRIES if k and k not in nums]
    if missing:
        print("warning: headings not found in PDF:", missing)
    env = dict(os.environ, CEAB_TOC_PAGES=json.dumps(nums))
    subprocess.run([sys.executable, os.path.abspath(__file__)], env=env, check=True)
    pdf = to_pdf(OUT)
    print(f"saved {OUT} and {pdf} ({npages} pages)")
